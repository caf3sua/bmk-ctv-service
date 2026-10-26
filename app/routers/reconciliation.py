import math
import os
import re
from datetime import datetime, timezone, date, timedelta
from dateutil.relativedelta import relativedelta
from io import BytesIO
from typing import List, Optional
from urllib.parse import quote
from bson import ObjectId
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pymongo import UpdateOne

from app.core.config import settings
from app.core.database import get_db
from app.core.logging import get_logger
from app.core.s3 import get_s3_client, upload_to_s3
from app.core.security import get_current_user
from app.models.reconciliation import (
    ReconciliationRecordCreate,
    ReconciliationRecordUpdate,
    ReconciliationRecordResponse,
    ReconciliationListResponse,
    ReconciliationHistoryResponse,
    ReconciliationHistoryListResponse,
    ReconciliationHistoryStats,
    ReconciliationResultFileInfo,
    ImportHrTpBankResult,
    ExpiringContractsStatsResponse,
    ImportExpiringContractsResult,
    ExpiringContractsHistoryResponse,
    ExpiringContractsHistoryListResponse,
    BmkSystemInfo,
    TpBankInfo,
    BmkHrInfo,
    CalculateContractExpiryRequest,
    CalculateContractExpiryResponse,
)

SERVICE_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DAO_HAN_TEMPLATE_PATH = os.path.join(SERVICE_ROOT, "templates", "template_bmk_ngay_dao_han.xlsx")
EXPORT_DAO_HAN_TEMPLATE_PATH = os.path.join(SERVICE_ROOT, "templates", "template_export_danhsach_ctv_dao_han.xlsx")

router = APIRouter(prefix="/api/reconciliation", tags=["Reconciliation"])
logger = get_logger(__name__)

COLLECTION = "bmk_ctv_reconciliations"
COLLABORATORS_COLLECTION = "bmk_ctv_collaborators"
HISTORY_COLLECTION = "bmk_ctv_reconciliation_history"
EXPIRING_HISTORY_COLLECTION = "bmk_ctv_expiring_contracts_history"

def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

def _to_response(doc: dict) -> dict:
    created_source = doc.get("createdSource") or doc.get("source") or "bmk_system"
    tpbank_info = doc.get("tpbankInfo") or {"contracts": []}
    if isinstance(tpbank_info, dict):
        tpbank_info.setdefault("contracts", [])

    return {
        "id": str(doc["_id"]),
        "employeeCode": doc.get("employeeCode", ""),
        "fullName": doc.get("fullName", ""),
        "idNumber": doc.get("idNumber") or doc.get("cccd"),
        "createdSource": created_source,
        "isBmkSystemExist": doc.get("isBmkSystemExist", False),
        "isSynced": doc.get("isSynced", False),
        "result": doc.get("result"),
        "departmentLevel1": doc.get("departmentLevel1"),
        "position": doc.get("position"),
        "employmentStatus": doc.get("employmentStatus"),
        "onboardDate": doc.get("onboardDate"),
        "offboardDate": doc.get("offboardDate"),
        "calContractExpiryDate": doc.get("calContractExpiryDate"),
        "tpbankInfo": tpbank_info,
        "bmkHrInfo": doc.get("bmkHrInfo") or {
            "contractCount": 0,
            "idCardCount": 0,
            "liquidationCount": 0,
            "taxCommitmentCount": 0,
        },
        "bmkSystemInfo": doc.get("bmkSystemInfo") or {
            "contractCount": 0,
            "idCardCount": 0,
            "liquidationCount": 0,
            "taxCommitmentCount": 0,
        },
        "reconciliationStatus": doc.get("reconciliationStatus", "pending"),
        "createdAt": doc.get("createdAt", _now()),
        "updatedAt": doc.get("updatedAt", _now()),
    }


def _parse_excel_date(val) -> tuple[Optional[str], Optional[datetime]]:
    if val is None:
        return None, None
    if isinstance(val, (datetime, date)):
        if val.year == 1900:
            return None, None
        dt = datetime(val.year, val.month, val.day) if isinstance(val, date) and not isinstance(val, datetime) else val
        return dt.strftime("%Y-%m-%d"), dt
    
    s = str(val).strip()
    if not s or s.upper() in ["N/A", "NONE", "-", "NULL"] or "1900" in s:
        return None, None
    
    for fmt in ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"]:
        try:
            dt = datetime.strptime(s, fmt)
            if dt.year == 1900:
                return None, None
            return dt.strftime("%Y-%m-%d"), dt
        except ValueError:
            pass
    return None, None


def _calculate_system_info(collab_doc: dict) -> dict:
    checklist = collab_doc.get("checklist") or {}
    
    # CCCD: 1 nếu cccd.checked = true hoặc cccd.file khác null
    cccd = checklist.get("cccd") or {}
    id_card_count = 1 if (cccd.get("checked") or cccd.get("file")) else 0
    
    # Cam kết thuế: 1 nếu ckt.checked = true hoặc ckt.file khác null
    ckt = checklist.get("ckt") or {}
    tax_count = 1 if (ckt.get("checked") or ckt.get("file")) else 0
    
    # Biên bản thanh lý: 1 nếu bbtl.date != null hoặc bbtl.file khác null
    bbtl = checklist.get("bbtl") or {}
    liquidation_count = 1 if (bbtl.get("date") or bbtl.get("file")) else 0
    
    # Hợp đồng dịch vụ: count từ hddv.contract_date (khởi tạo với startDate & endDate khác null)
    hddv = checklist.get("hddv") or {}
    contract_dates = hddv.get("contract_date") or []
    contract_count = 0
    for d in contract_dates:
        if isinstance(d, dict):
            start_d = d.get("startDate")
            end_d = d.get("endDate")
            if start_d and end_d:
                contract_count += 1

    return {
        "contractCount": contract_count,
        "idCardCount": id_card_count,
        "liquidationCount": liquidation_count,
        "taxCommitmentCount": tax_count,
    }

@router.get("", response_model=ReconciliationListResponse)
async def list_reconciliation_records(
    keyword: Optional[str] = Query(None, description="Tìm theo mã CTV, họ tên, đơn vị"),
    employment_status: Optional[str] = Query(None, description="Lọc theo tình trạng nhân sự"),
    reconciliation_status: Optional[str] = Query(None, description="Lọc theo trạng thái đối soát"),
    created_source: Optional[str] = Query(None, description="Lọc theo nguồn tạo (bmk_system, bmk_hr, tpbank)"),
    is_synced: Optional[str] = Query(None, description="Lọc theo trạng thái đối soát (synced, unsynced)"),
    result_status: Optional[str] = Query(None, description="Lọc theo kết quả đối soát (all, match_all, mismatch_contract, mismatch_idcard, mismatch_liquidation)"),
    page: int = Query(1, ge=1, description="Số trang"),
    page_size: int = Query(20, ge=1, le=200, description="Số dòng mỗi trang"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    conditions = []
    if keyword and keyword.strip():
        term = keyword.strip()
        conditions.append({
            "$or": [
                {"employeeCode": {"$regex": re.escape(term), "$options": "i"}},
                {"fullName": {"$regex": re.escape(term), "$options": "i"}},
                {"idNumber": {"$regex": re.escape(term), "$options": "i"}},
                {"departmentLevel1": {"$regex": re.escape(term), "$options": "i"}},
                {"position": {"$regex": re.escape(term), "$options": "i"}},
            ]
        })

    if employment_status and employment_status.strip() and employment_status.strip() != "all":
        conditions.append({"employmentStatus": employment_status.strip()})

    if reconciliation_status and reconciliation_status.strip() and reconciliation_status.strip() != "all":
        conditions.append({"reconciliationStatus": reconciliation_status.strip()})

    if is_synced and is_synced.strip() and is_synced.strip() != "all":
        val = is_synced.strip().lower()
        if val in ["synced", "true", "1"]:
            conditions.append({"isSynced": True})
        elif val in ["unsynced", "false", "0"]:
            conditions.append({
                "$or": [
                    {"isSynced": False},
                    {"isSynced": {"$exists": False}},
                ]
            })

    if result_status and result_status.strip() and result_status.strip() != "all":
        r_val = result_status.strip().lower()
        if r_val in ["match_all", "success_all"]:
            conditions.append({
                "result.contract": "success",
                "result.idCard": "success",
                "result.liquidation": "success",
            })
        elif r_val in ["warn_bank_contract", "warn_contract", "warn", "mismatch_bank"]:
            conditions.append({"result.contract": "warn"})
        elif r_val in ["mismatch_contract", "failed_contract"]:
            conditions.append({"result.contract": "failed"})
        elif r_val in ["mismatch_idcard", "failed_idcard"]:
            conditions.append({"result.idCard": "failed"})
        elif r_val in ["mismatch_liquidation", "failed_liquidation"]:
            conditions.append({"result.liquidation": "failed"})

    if created_source and created_source.strip() and created_source.strip() != "all":
        src = created_source.strip()
        if src == "bmk_system":
            conditions.append({
                "$or": [
                    {"createdSource": "bmk_system"},
                    {"createdSource": "BMK System"},
                    {"createdSource": {"$exists": False}},
                ]
            })
        elif src == "bmk_hr":
            conditions.append({
                "$or": [
                    {"createdSource": "bmk_hr"},
                    {"createdSource": "HR BMK"},
                ]
            })
        elif src in ["hr_tpbank", "HR TP Bank", "HR Tp bank"]:
            conditions.append({
                "$or": [
                    {"createdSource": "hr_tpbank"},
                    {"createdSource": "HR TP Bank"},
                    {"createdSource": "HR Tp bank"},
                ]
            })
        elif src == "tpbank":
            conditions.append({
                "$or": [
                    {"createdSource": "tpbank"},
                    {"createdSource": "TP Bank"},
                ]
            })
        else:
            conditions.append({"createdSource": src})

    query = {"$and": conditions} if conditions else {}

    total = await db[COLLECTION].count_documents(query)
    total_pages = max(1, math.ceil(total / page_size))
    skip = (page - 1) * page_size

    cursor = db[COLLECTION].find(query).sort("updatedAt", -1).skip(skip).limit(page_size)
    items = []
    async for doc in cursor:
        items.append(_to_response(doc))

    return {
        "items": items,
        "total": total,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
    }


@router.get("/history", response_model=ReconciliationHistoryListResponse)
async def list_reconciliation_history(
    page: int = Query(1, ge=1, description="Số trang"),
    page_size: int = Query(20, ge=1, le=100, description="Số dòng mỗi trang"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách lịch sử các phiên đối soát TP Bank."""
    total = await db[HISTORY_COLLECTION].count_documents({})
    total_pages = max(1, math.ceil(total / page_size))
    skip = (page - 1) * page_size

    cursor = db[HISTORY_COLLECTION].find({}).sort("createdAt", -1).skip(skip).limit(page_size)
    items = []
    async for doc in cursor:
        items.append({
            "id": str(doc["_id"]),
            "filename": doc.get("filename", ""),
            "s3Key": doc.get("s3Key", ""),
            "s3Bucket": doc.get("s3Bucket", settings.S3_BUCKET),
            "fileSize": doc.get("fileSize", 0),
            "uploadedBy": doc.get("uploadedBy", ""),
            "username": doc.get("username", ""),
            "totalRows": doc.get("totalRows", 0),
            "successRows": doc.get("successRows", 0),
            "failedRows": doc.get("failedRows", 0),
            "status": doc.get("status", "success"),
            "message": doc.get("message", ""),
            "stats": doc.get("stats") or {
                "totalSuccess": 0,
                "totalWarnBank": 0,
                "totalMismatchContract": 0,
                "totalMismatchIdCard": 0,
                "totalMismatchLiquidation": 0,
            },
            "resultFile": doc.get("resultFile"),
            "createdAt": doc.get("createdAt", _now()),
        })

    return {
        "items": items,
        "total": total,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
    }


@router.get("/history/{history_id}/download")
async def download_reconciliation_history_file(
    history_id: str,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Tải file đối soát gốc đã upload từ S3 MinIO."""
    if not ObjectId.is_valid(history_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mã lịch sử đối soát không hợp lệ",
        )

    history = await db[HISTORY_COLLECTION].find_one({"_id": ObjectId(history_id)})
    if not history:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy lịch sử đối soát",
        )

    s3_key = history.get("s3Key")
    s3_bucket = history.get("s3Bucket", settings.S3_BUCKET)
    filename = history.get("filename", "reconciliation_source.xlsx")

    if not s3_key:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không có thông tin lưu trữ S3 cho file này",
        )

    try:
        s3 = get_s3_client()
        response = s3.get_object(Bucket=s3_bucket, Key=s3_key)

        def iter_chunks():
            for chunk in response["Body"].iter_chunks(chunk_size=1024 * 1024):
                yield chunk

        return StreamingResponse(
            iter_chunks(),
            media_type=response.get("ContentType", "application/octet-stream"),
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
            },
        )
    except Exception as e:
        logger.error(f"Lỗi khi tải file từ S3 cho History ID={history_id}: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể tải file từ S3 MinIO",
        )


@router.get("/history/{history_id}/download-result")
async def download_reconciliation_result_file(
    history_id: str,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Tải file kết quả đối soát đã bổ sung các cột đối soát từ S3 MinIO."""
    if not ObjectId.is_valid(history_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mã lịch sử đối soát không hợp lệ",
        )

    history = await db[HISTORY_COLLECTION].find_one({"_id": ObjectId(history_id)})
    if not history:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy lịch sử đối soát",
        )

    result_file = history.get("resultFile")
    if not result_file or not isinstance(result_file, dict) or not result_file.get("s3Key"):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Phiên đối soát này chưa có file kết quả",
        )

    s3_key = result_file.get("s3Key")
    s3_bucket = result_file.get("s3Bucket", settings.S3_BUCKET)
    filename = result_file.get("filename", "ket_qua_doi_soat.xlsx")

    try:
        s3 = get_s3_client()
        response = s3.get_object(Bucket=s3_bucket, Key=s3_key)

        def iter_chunks():
            for chunk in response["Body"].iter_chunks(chunk_size=1024 * 1024):
                yield chunk

        return StreamingResponse(
            iter_chunks(),
            media_type=response.get("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
            },
        )
    except Exception as e:
        logger.error(f"Lỗi khi tải file kết quả từ S3 cho History ID={history_id}: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể tải file kết quả từ S3 MinIO",
        )



@router.get("/expiring-contracts/stats", response_model=ExpiringContractsStatsResponse)
async def get_expiring_contracts_stats(
    days: int = Query(30, ge=1, le=180, description="Số ngày tính movement data"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Thống kê cho page Hợp đồng đến hạn:
    - Pie chart: số lượng CTV nghỉ việc vs còn hiệu lực
    - Bar chart: số lượng hợp đồng đến hạn theo từng ngày (movement data)
    """
    # 1. Pie chart stats
    active_count = await db[COLLECTION].count_documents({
        "employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]}
    })
    resigned_count = await db[COLLECTION].count_documents({
        "employmentStatus": "Nghỉ việc"
    })
    total = active_count + resigned_count

    # Thống kê CTV còn hiệu lực có ngày đáo hạn vs thiếu ngày đáo hạn
    active_with_expiry = await db[COLLECTION].count_documents({
        "employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]},
        "calContractExpiryDate": {"$nin": [None, "", " ", "-"]}
    })
    active_without_expiry = max(0, active_count - active_with_expiry)

    # 2. Bar chart movement data trong N ngày tới
    today = datetime.now(timezone.utc).date()
    start_date_str = today.strftime("%Y-%m-%d")
    end_date_str = (today + timedelta(days=days - 1)).strftime("%Y-%m-%d")

    pipeline = [
        {
            "$match": {
                "employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]},
                "calContractExpiryDate": {
                    "$gte": start_date_str,
                    "$lte": end_date_str
                }
            }
        },
        {
            "$group": {
                "_id": "$calContractExpiryDate",
                "count": {"$sum": 1}
            }
        }
    ]

    counts_by_date = {}
    async for row in db[COLLECTION].aggregate(pipeline):
        d_val = row["_id"]
        if d_val:
            counts_by_date[str(d_val)] = row["count"]

    movement_items = []
    total_expiring = 0
    for i in range(days):
        cur_date = today + timedelta(days=i)
        cur_date_str = cur_date.strftime("%Y-%m-%d")
        cnt = counts_by_date.get(cur_date_str, 0)
        total_expiring += cnt
        movement_items.append({
            "date": cur_date_str,
            "label": cur_date.strftime("%d/%m"),
            "count": cnt,
        })

    return {
        "pieChart": {
            "active": active_count,
            "resigned": resigned_count,
            "total": total,
            "activeWithExpiry": active_with_expiry,
            "activeWithoutExpiry": active_without_expiry,
        },
        "barChart": {
            "days": days,
            "totalExpiring": total_expiring,
            "items": movement_items,
        }
    }


@router.get("/expiring-contracts", response_model=ReconciliationListResponse)
async def list_expiring_contracts(
    keyword: Optional[str] = Query(None, description="Tìm theo mã CTV, họ tên, số CCCD, đơn vị"),
    days: Optional[int] = Query(None, ge=1, le=180, description="Lọc hợp đồng đến hạn trong N ngày tới"),
    is_calculate_expired: Optional[bool] = Query(None, description="Lọc CTV có trong đợt upload tính ngày đáo hạn vừa rồi"),
    page: int = Query(1, ge=1, description="Số trang"),
    page_size: int = Query(20, ge=1, le=200, description="Số dòng mỗi trang"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách các hợp đồng đến hạn từ collection bmk_ctv_reconciliations
    với employee còn hiệu lực, sắp xếp theo contractExpiryDate tăng dần.
    """
    conditions = [
        {"employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]}}
    ]

    if is_calculate_expired is True:
        conditions.append({"isCalculateExpired": True})
    elif is_calculate_expired is False:
        conditions.append({"$or": [{"isCalculateExpired": None}, {"isCalculateExpired": False}]})

    if keyword and keyword.strip():
        term = keyword.strip()
        conditions.append({
            "$or": [
                {"employeeCode": {"$regex": re.escape(term), "$options": "i"}},
                {"fullName": {"$regex": re.escape(term), "$options": "i"}},
                {"idNumber": {"$regex": re.escape(term), "$options": "i"}},
                {"departmentLevel1": {"$regex": re.escape(term), "$options": "i"}},
            ]
        })

    if days is not None and days > 0:
        today_str = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        end_date_str = (datetime.now(timezone.utc) + timedelta(days=days - 1)).strftime("%Y-%m-%d")
        conditions.append({
            "calContractExpiryDate": {
                "$gte": today_str,
                "$lte": end_date_str
            }
        })

    query = {"$and": conditions} if conditions else {}

    total = await db[COLLECTION].count_documents(query)
    total_pages = max(1, math.ceil(total / page_size))
    skip = (page - 1) * page_size

    # Sắp xếp theo calContractExpiryDate tăng dần:
    # Các bản ghi có calContractExpiryDate (hợp lệ) lên đầu, tiếp đến các bản ghi chưa có ngày đáo hạn
    pipeline = [
        {"$match": query},
        {
            "$addFields": {
                "_hasExpiry": {
                    "$cond": [
                        {
                            "$and": [
                                {"$eq": [{"$type": "$calContractExpiryDate"}, "string"]},
                                {"$gt": [{"$strLenCP": "$calContractExpiryDate"}, 5]}
                            ]
                        },
                        0,
                        1
                    ]
                }
            }
        },
        {"$sort": {"_hasExpiry": 1, "calContractExpiryDate": 1, "employeeCode": 1}},
        {"$skip": skip},
        {"$limit": page_size},
    ]

    items = []
    async for doc in db[COLLECTION].aggregate(pipeline):
        items.append(_to_response(doc))

    return {
        "items": items,
        "total": total,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
    }


def _get_notification_text(expiry_date_str: Optional[str]) -> str:
    if not expiry_date_str or not str(expiry_date_str).strip():
        return "Thiếu ngày đáo hạn"
    try:
        exp = datetime.strptime(str(expiry_date_str)[:10], "%Y-%m-%d").date()
        today = datetime.now(timezone.utc).date()
        diff_days = (exp - today).days
        if diff_days < 0:
            return f"Hợp đồng đã quá hạn {abs(diff_days)} ngày"
        else:
            return f"Hợp đồng sẽ đáo hạn trong {diff_days} ngày"
    except Exception:
        return "Thiếu ngày đáo hạn"


def _to_excel_date(d_val):
    if not d_val:
        return ""
    if isinstance(d_val, (datetime, date)):
        return d_val if isinstance(d_val, date) else d_val.date()
    s = str(d_val).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return s


def _generate_expiring_excel_file(items: List[dict], title_month: Optional[str] = None) -> bytes:
    wb_res = load_workbook(DAO_HAN_TEMPLATE_PATH)
    ws_res = wb_res.active

    if title_month:
        orig_title = ws_res.cell(row=1, column=2).value or "DANH SÁCH CẬP NHẬT NGÀY ĐÁO HẠN HỢP ĐỒNG"
        ws_res.cell(row=1, column=2, value=f"{orig_title} - THÁNG {title_month}")

    # Cột F (index 6): Ngày đáo hạn dự kiến
    ws_res.cell(row=3, column=6, value="Ngày đáo hạn dự kiến")

    # Cột G (index 7): Thông báo
    cell_g3 = ws_res.cell(row=3, column=7, value="Thông báo")
    cell_f3 = ws_res.cell(row=3, column=6)
    if cell_f3.font:
        cell_g3.font = cell_f3.font.copy()
    if cell_f3.alignment:
        cell_g3.alignment = cell_f3.alignment.copy()
    if cell_f3.border:
        cell_g3.border = cell_f3.border.copy()
    if cell_f3.fill:
        cell_g3.fill = cell_f3.fill.copy()

    # Xóa các dòng mẫu cũ bên dưới header (từ dòng 4 trở đi)
    if ws_res.max_row >= 4:
        ws_res.delete_rows(4, ws_res.max_row - 3)

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    data_font = Font(name="Tahoma", size=10)
    overdue_font = Font(name="Tahoma", size=10, color="DC2626", bold=True)

    for idx, it in enumerate(items, start=1):
        r_num = 3 + idx
        c_stt = ws_res.cell(row=r_num, column=1, value=idx)
        c_code = ws_res.cell(row=r_num, column=2, value=it.get("employeeCode"))
        c_name = ws_res.cell(row=r_num, column=3, value=it.get("fullName"))
        c_cccd = ws_res.cell(row=r_num, column=4, value=it.get("idNumber"))

        onboard_val = _to_excel_date(it.get("onboardDate"))
        c_onboard = ws_res.cell(row=r_num, column=5, value=onboard_val)
        if isinstance(onboard_val, (date, datetime)):
            c_onboard.number_format = "dd/mm/yyyy"

        exp_val = _to_excel_date(it.get("calContractExpiryDate"))
        c_exp = ws_res.cell(row=r_num, column=6, value=exp_val)
        if isinstance(exp_val, (date, datetime)):
            c_exp.number_format = "dd/mm/yyyy"

        notif_str = it.get("notificationText") or _get_notification_text(it.get("calContractExpiryDate"))
        c_notif = ws_res.cell(row=r_num, column=7, value=notif_str)

        c_stt.alignment = align_center
        c_code.alignment = align_center
        c_name.alignment = align_left
        c_cccd.alignment = align_center
        c_onboard.alignment = align_center
        c_exp.alignment = align_center
        c_notif.alignment = align_left

        for c in (c_stt, c_code, c_name, c_cccd, c_onboard, c_exp):
            c.font = data_font
            c.border = thin_border

        c_notif.border = thin_border
        if "quá hạn" in str(notif_str).lower():
            c_notif.font = overdue_font
        else:
            c_notif.font = data_font

    ws_res.cell(row=2, column=3, value=len(items))
    ws_res.column_dimensions["G"].width = 38

    output = BytesIO()
    wb_res.save(output)
    wb_res.close()
    return output.getvalue()


def _generate_expiring_export_file(items: List[dict], target_year: int, target_month: int) -> bytes:
    if not os.path.exists(EXPORT_DAO_HAN_TEMPLATE_PATH):
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không tìm thấy file template mẫu template_export_danhsach_ctv_dao_han.xlsx",
        )

    wb = load_workbook(EXPORT_DAO_HAN_TEMPLATE_PATH)
    ws = wb.active

    # Cập nhật tên sheet theo tháng (ví dụ '10 2026')
    sheet_title = f"{target_month:02d} {target_year}"
    ws.title = sheet_title

    # Cập nhật tiêu đề ở dòng 4
    ws.cell(row=4, column=1, value=f"DANH SÁCH CỘNG TÁC VIÊN ĐÁO HẠN THÁNG {target_month:02d}/{target_year}")

    # Xóa các dòng mẫu cũ từ dòng 7 trở đi nếu có
    if ws.max_row >= 7:
        ws.delete_rows(7, ws.max_row - 6)

    thin_side = Side(style="thin", color="000000")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    font_data = Font(name="Times New Roman", size=11)

    for idx, it in enumerate(items, start=1):
        r_num = 6 + idx

        # 1. STT (Col A)
        c_stt = ws.cell(row=r_num, column=1, value=idx)
        c_stt.alignment = align_center

        # 2. Ngày BMK thanh toán (Col B)
        pay_date_val = _to_excel_date(it.get("paymentDate"))
        c_pay = ws.cell(row=r_num, column=2, value=pay_date_val if pay_date_val else None)
        c_pay.alignment = align_center
        if isinstance(pay_date_val, (date, datetime)):
            c_pay.number_format = "dd/mm/yyyy"

        # 3. MNV (Col C)
        code_val = it.get("employeeCode", "")
        c_code = ws.cell(row=r_num, column=3, value=str(code_val).strip() if code_val else "")
        c_code.alignment = align_center
        c_code.number_format = "@"

        # 4. HỌ VÀ TÊN (Col D)
        c_name = ws.cell(row=r_num, column=4, value=it.get("fullName", ""))
        c_name.alignment = align_left

        # 5. MÃ SỐ THUẾ (Col E)
        tax_val = it.get("taxCode") or it.get("idNumber") or ""
        c_tax = ws.cell(row=r_num, column=5, value=str(tax_val).strip() if tax_val else "")
        c_tax.alignment = align_center
        c_tax.number_format = "@"

        # 6. THU NHẬP CHỊU THUẾ (Col F)
        c_income = ws.cell(row=r_num, column=6, value=None)
        c_income.alignment = align_right
        c_income.number_format = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'

        # 7. THUẾ TNCN (Col G)
        c_tax_amount = ws.cell(row=r_num, column=7, value=None)
        c_tax_amount.alignment = align_right
        c_tax_amount.number_format = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'

        # 8. THỰC LĨNH (Col H)
        c_net = ws.cell(row=r_num, column=8, value=None)
        c_net.alignment = align_right
        c_net.number_format = '_(* #,##0_);_(* (#,##0);_(* "-"??_);_(@_)'

        # 9. NGÀY GIA NHẬP (Col I)
        onboard_val = _to_excel_date(it.get("onboardDate"))
        c_onboard = ws.cell(row=r_num, column=9, value=onboard_val if onboard_val else None)
        c_onboard.alignment = align_center
        if isinstance(onboard_val, (date, datetime)):
            c_onboard.number_format = "dd/mm/yyyy"

        # 10. NGÀY THÔI VIỆC (Col J)
        offboard_val = _to_excel_date(it.get("offboardDate"))
        c_offboard = ws.cell(row=r_num, column=10, value=offboard_val if offboard_val else None)
        c_offboard.alignment = align_center
        if isinstance(offboard_val, (date, datetime)):
            c_offboard.number_format = "dd/mm/yyyy"

        for col_idx in range(1, 11):
            cell = ws.cell(row=r_num, column=col_idx)
            cell.font = font_data
            cell.border = thin_border

    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


@router.post("/calculate-contract-expiry", response_model=CalculateContractExpiryResponse)
async def calculate_contract_expiry(
    payload: CalculateContractExpiryRequest,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Tính ngày đáo hạn dự kiến cho các CTV còn hiệu lực theo tháng được chọn:
    - Với mỗi cộng tác viên còn hiệu lực, từ ngày gia nhập tính ngày đáo hạn dự kiến (cộng thêm 6 tháng):
      + Nếu ngày đáo hạn dự kiến có tháng nhỏ hơn Selected month thì tiếp tục cộng thêm chu kỳ 6 tháng.
      + Nếu ngày đáo hạn dự kiến có tháng >= selected month thì dừng lại.
    - Cập nhật calContractExpiryDate vào collection bmk_ctv_reconciliations.
    """
    month_str = payload.month.strip()
    try:
        parts = month_str.split("-")
        target_year = int(parts[0])
        target_month = int(parts[1])
        if not (1 <= target_month <= 12):
            raise ValueError()
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Định dạng tháng không hợp lệ (cần 'YYYY-MM', ví dụ '2026-10'): {month_str}",
        )

    # Lấy danh sách CTV còn hiệu lực
    active_cursor = db[COLLECTION].find(
        {"employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]}},
        {"_id": 1, "onboardDate": 1}
    )

    operations = []
    now = _now()
    calculated_count = 0
    total_active = 0

    async for doc in active_cursor:
        total_active += 1
        raw_onboard = doc.get("onboardDate")
        onboard_str, onboard_dt = _parse_excel_date(raw_onboard)

        if not onboard_dt:
            continue

        # Quy tắc đối soát: Nếu ngày gia nhập trước ngày 1/4/2025 thì Ngày gia nhập = 1/4/2025
        if onboard_dt < datetime(2025, 4, 1):
            onboard_dt = datetime(2025, 4, 1)

        # Tính chu kỳ 6 tháng:
        # + Nếu ngày đáo hạn dự kiến có tháng nhỏ hơn Selected month thì tiếp tục cộng thêm chu kỳ 6 tháng.
        # + Nếu ngày đáo hạn dự kiến có tháng >= selected month thì dừng lại
        expiry_dt = onboard_dt + relativedelta(months=6)
        while (expiry_dt.year < target_year) or (expiry_dt.year == target_year and expiry_dt.month < target_month):
            expiry_dt = expiry_dt + relativedelta(months=6)

        cal_expiry_str = expiry_dt.strftime("%Y-%m-%d")
        calculated_count += 1

        operations.append(
            UpdateOne(
                {"_id": doc["_id"]},
                {
                    "$set": {
                        "calContractExpiryDate": cal_expiry_str,
                        "updatedAt": now,
                    }
                }
            )
        )

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    return CalculateContractExpiryResponse(
        status="success",
        message=f"Đã tính ngày đáo hạn dự kiến cho {calculated_count} CTV còn hiệu lực theo tháng {target_month:02d}/{target_year}.",
        calculatedCount=calculated_count,
        totalActive=total_active,
        month=f"{target_year}-{target_month:02d}",
    )


@router.get("/export-expiring-contracts")
async def export_expiring_contracts(
    month: str = Query(..., description="Tháng cần xuất dữ liệu theo định dạng YYYY-MM (ví dụ 2026-10)"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Xuất danh sách nhân viên còn hiệu lực có tháng của ngày đáo hạn dự kiến = selected month theo template_export_danhsach_ctv_dao_han.xlsx."""
    month_str = month.strip()
    try:
        parts = month_str.split("-")
        target_year = int(parts[0])
        target_month = int(parts[1])
        if not (1 <= target_month <= 12):
            raise ValueError()
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Định dạng tháng không hợp lệ (cần 'YYYY-MM', ví dụ '2026-10'): {month_str}",
        )

    month_prefix = f"{target_year}-{target_month:02d}"

    # Lấy tất cả data của nhân viên còn hiệu lực và tháng của ngày đáo hạn dự kiến = selected month
    cursor = db[COLLECTION].find({
        "employmentStatus": {"$in": ["Hiệu lực", "Hiện diện"]},
        "calContractExpiryDate": {"$regex": f"^{month_prefix}"}
    }).sort([("calContractExpiryDate", 1), ("employeeCode", 1)])

    docs = []
    emp_codes = []
    async for doc in cursor:
        docs.append(doc)
        code = doc.get("employeeCode")
        if code:
            emp_codes.append(str(code).strip())

    # Map taxCode từ collaborators nếu có
    tax_code_map = {}
    if emp_codes:
        collab_cursor = db[COLLABORATORS_COLLECTION].find(
            {"employeeCode": {"$in": emp_codes}},
            {"employeeCode": 1, "taxCode": 1, "idNumber": 1}
        )
        async for c_doc in collab_cursor:
            c_code = c_doc.get("employeeCode")
            if c_code:
                tax_code_map[str(c_code).strip()] = c_doc.get("taxCode") or c_doc.get("idNumber")

    items = []
    for doc in docs:
        emp_code = str(doc.get("employeeCode", "")).strip()
        tax_code = tax_code_map.get(emp_code) or doc.get("taxCode") or doc.get("idNumber") or doc.get("cccd") or ""
        items.append({
            "employeeCode": emp_code,
            "fullName": doc.get("fullName", ""),
            "taxCode": tax_code,
            "idNumber": doc.get("idNumber") or doc.get("cccd") or "",
            "onboardDate": doc.get("onboardDate") or "",
            "offboardDate": doc.get("offboardDate") or None,
            "paymentDate": doc.get("paymentDate") or None,
            "calContractExpiryDate": doc.get("calContractExpiryDate") or "",
            "notificationText": _get_notification_text(doc.get("calContractExpiryDate")),
        })

    file_bytes = _generate_expiring_export_file(items, target_year=target_year, target_month=target_month)

    filename = f"danh_sach_ctv_dao_han_thang_{target_month:02d}_{target_year}.xlsx"
    encoded_filename = quote(filename)
    headers = {
        "Content-Disposition": f"attachment; filename=\"{filename}\"; filename*=UTF-8''{encoded_filename}",
        "Access-Control-Expose-Headers": "Content-Disposition",
    }
    return StreamingResponse(
        BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/template-expiring-contracts")
async def download_expiring_contracts_template(
    current_user: dict = Depends(get_current_user),
):
    """Download template mẫu template_bmk_ngay_dao_han.xlsx để import ngày đáo hạn."""
    if not os.path.exists(DAO_HAN_TEMPLATE_PATH):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy file template mẫu template_bmk_ngay_dao_han.xlsx",
        )
    return FileResponse(
        DAO_HAN_TEMPLATE_PATH,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="template_bmk_ngay_dao_han.xlsx",
    )


@router.post("/import-expiring-contracts", response_model=ImportExpiringContractsResult)
async def import_expiring_contracts_data(
    file: UploadFile = File(...),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Import data ngày gia nhập từ file Excel:
    - Đọc dữ liệu cột E (Ngày gia nhập) lưu vào onboardDate ở collection bmk_ctv_reconciliations theo quy tắc file đối soát.
    - Giá trị field calContractExpiryDate không tính khi import (sẽ tính ở logic riêng).
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file Excel (.xlsx, .xlsm)",
        )

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Không đọc được file Excel: {str(e)}",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File Excel không chứa dữ liệu",
        )

    header_row_idx = None
    header = []
    for idx, r in enumerate(rows):
        if r and any(
            cell and any(k in str(cell).upper() for k in ["MNV", "MÃ NV", "MÃ NHÂN VIÊN", "MÃ CTV", "EMPLOYEE CODE"])
            for cell in r
        ):
            header_row_idx = idx
            header = [str(c).strip() if c is not None else "" for c in r]
            break

    if header_row_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không tìm thấy dòng tiêu đề chứa cột 'MNV' (Mã nhân viên) trong file Excel",
        )

    header_upper = [h.upper() for h in header]

    def find_col_idx(candidates: List[str]) -> Optional[int]:
        for candidate in candidates:
            cand_u = candidate.upper()
            for idx, h in enumerate(header_upper):
                if cand_u in h:
                    return idx
        return None

    code_idx = find_col_idx(["MNV", "MÃ NV", "MÃ NHÂN VIÊN", "MÃ CTV", "EMPLOYEE CODE", "MA NV"])
    name_idx = find_col_idx(["HỌ VÀ TÊN", "HỌ TÊN", "HO VA TEN", "FULL NAME", "NAME"])
    cccd_idx = find_col_idx(["SỐ CCCD", "CCCD", "CMND", "SỐ CMND", "ID NUMBER", "SO CCCD"])
    onboard_idx = find_col_idx([
        "NGÀY GIA NHẬP", "GIA NHẬP", "NGAY GIA NHAP", "GIA NHAP",
        "NGÀY VÀO", "NGAY VAO", "ONBOARD DATE", "ONBOARD", "JOIN DATE"
    ])
    expiry_idx = find_col_idx(["NGÀY ĐÁO HẠN", "ĐÁO HẠN", "NGAY DAO HAN", "DAO HAN", "NGÀY HẾT HẠN", "HẾT HẠN", "NGÀY ĐẾN HẠN", "ĐẾN HẠN", "EXPIRY"])

    if code_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File thiếu cột 'MNV' (Mã nhân viên)",
        )

    if onboard_idx is None:
        # Fallback sang cột E (index 4) nếu file có ít nhất 5 cột
        if len(header) > 4:
            onboard_idx = 4
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File thiếu cột 'Ngày gia nhập' (Cột E)",
            )

    def parse_id_number(val) -> Optional[str]:
        if val is None:
            return None
        s = str(val).strip()
        if not s or s == "-":
            return None
        if s.endswith(".0"):
            s = s[:-2]
        return s

    # Helpers _get_notification_text and _generate_expiring_excel_file are defined at module level

    operations = []
    now = _now()
    created_count = 0
    updated_count = 0
    total_processed = 0
    processed_items = []

    existing_codes = set()
    async for doc in db[COLLECTION].find({}, {"employeeCode": 1}):
        c = doc.get("employeeCode")
        if c:
            existing_codes.add(str(c).strip())

    for row in rows[header_row_idx + 1 :]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_code = row[code_idx] if code_idx < len(row) else None
        if raw_code is None:
            continue
        employee_code = str(raw_code).strip()
        if not employee_code:
            continue

        if employee_code.isdigit() and len(employee_code) < 5:
            employee_code = employee_code.zfill(5)

        full_name = str(row[name_idx]).strip() if (name_idx is not None and name_idx < len(row) and row[name_idx] is not None) else ""
        raw_cccd = row[cccd_idx] if (cccd_idx is not None and cccd_idx < len(row)) else None
        id_number = parse_id_number(raw_cccd)

        # Lấy giá trị cột E (Ngày gia nhập) lưu vào onboardDate theo rule đối soát
        raw_onboard = row[onboard_idx] if (onboard_idx is not None and onboard_idx < len(row)) else None
        onboard_str, onboard_dt = _parse_excel_date(raw_onboard)

        # Quy tắc đối soát: Nếu ngày gia nhập trước ngày 1/4/2025 (dd/MM/yyyy) thì Ngày gia nhập = 1/4/2025
        if onboard_dt and onboard_dt < datetime(2025, 4, 1):
            onboard_dt = datetime(2025, 4, 1)
            onboard_str = "2025-04-01"

        # calContractExpiryDate không tính khi import (sẽ tính ở logic riêng)
        notification_text = _get_notification_text(None)

        total_processed += 1
        processed_items.append({
            "employeeCode": employee_code,
            "fullName": full_name or "",
            "idNumber": id_number or "",
            "onboardDate": onboard_str or "",
            "calContractExpiryDate": "",
            "notificationText": notification_text,
        })

        if employee_code in existing_codes:
            # Nếu CTV đã tồn tại: chỉ cập nhật field onboardDate (không tính calContractExpiryDate khi import)
            updated_count += 1
            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {
                        "$set": {
                            "onboardDate": onboard_str,
                            "updatedAt": now,
                        }
                    }
                )
            )
        else:
            # Nếu CTV chưa tồn tại: tạo mới với onboardDate và calContractExpiryDate = None
            created_count += 1
            existing_codes.add(employee_code)
            new_doc = {
                "_id": employee_code,
                "employeeCode": employee_code,
                "fullName": full_name or "",
                "idNumber": id_number,
                "createdSource": "HTR BMK",
                "isBmkSystemExist": False,
                "isSynced": False,
                "departmentLevel1": None,
                "position": None,
                "employmentStatus": "Hiệu lực",
                "onboardDate": onboard_str,
                "offboardDate": None,
                "calContractExpiryDate": None,
                "tpbankInfo": {"contracts": []},
                "bmkHrInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "bmkSystemInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "result": None,
                "reconciliationStatus": "pending",
                "createdAt": now,
                "updatedAt": now,
            }
            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {"$setOnInsert": new_doc},
                    upsert=True,
                )
            )

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    # 1. Upload file gốc lên S3 (Requirement 4)
    now_ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    clean_name = re.sub(r"[^\w\.-]", "_", file.filename)
    s3_key = f"expiring_contracts/uploads/{now_ts}_{clean_name}"
    s3_bucket = settings.S3_BUCKET
    try:
        upload_to_s3(content, s3_key)
    except Exception as e:
        logger.error(f"Không thể upload file gốc ngày đáo hạn lên S3: {str(e)}")

    # 2. Tạo và upload file kết quả kèm cột Thông báo lên S3 (Requirement 4)
    result_file_info = None
    try:
        result_content = _generate_expiring_excel_file(processed_items)
        result_filename = f"ket_qua_ngay_dao_han_{clean_name}"
        result_s3_key = f"expiring_contracts/results/{now_ts}_{clean_name}"
        upload_to_s3(result_content, result_s3_key)
        result_file_info = {
            "filename": result_filename,
            "s3Key": result_s3_key,
            "s3Bucket": s3_bucket,
            "fileSize": len(result_content),
        }
    except Exception as e:
        logger.error(f"Không thể tạo hoặc upload file kết quả ngày đáo hạn lên S3: {str(e)}")

    # 3. Lưu lịch sử vào collection bmk_ctv_expiring_contracts_history (Requirement 4)
    uploaded_by = current_user.get("name") or current_user.get("fullName") or current_user.get("username", "Unknown")
    username = current_user.get("username", "Unknown")
    history_doc = {
        "filename": file.filename,
        "s3Key": s3_key,
        "s3Bucket": s3_bucket,
        "fileSize": len(content),
        "uploadedBy": uploaded_by,
        "username": username,
        "totalRows": total_processed,
        "createdCount": created_count,
        "updatedCount": updated_count,
        "status": "success",
        "message": f"Đã xử lý {total_processed} dòng: tạo mới {created_count} CTV, cập nhật {updated_count} CTV",
        "resultFile": result_file_info,
        "createdAt": now,
    }
    await db[EXPIRING_HISTORY_COLLECTION].insert_one(history_doc)

    return {
        "status": "success",
        "message": f"Đã xử lý {total_processed} dòng: cập nhật ngày gia nhập cho {updated_count} CTV (tạo mới {created_count} CTV).",
        "totalProcessed": total_processed,
        "createdCount": created_count,
        "updatedCount": updated_count,
    }


@router.get("/expiring-contracts/history", response_model=ExpiringContractsHistoryListResponse)
async def list_expiring_contracts_history(
    page: int = Query(1, ge=1, description="Số trang"),
    page_size: int = Query(20, ge=1, le=100, description="Số dòng mỗi trang"),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách lịch sử các phiên import ngày đáo hạn hợp đồng."""
    total = await db[EXPIRING_HISTORY_COLLECTION].count_documents({})
    total_pages = max(1, math.ceil(total / page_size))
    skip = (page - 1) * page_size

    cursor = db[EXPIRING_HISTORY_COLLECTION].find({}).sort("createdAt", -1).skip(skip).limit(page_size)
    items = []
    async for doc in cursor:
        items.append({
            "id": str(doc["_id"]),
            "filename": doc.get("filename", ""),
            "s3Key": doc.get("s3Key", ""),
            "s3Bucket": doc.get("s3Bucket", settings.S3_BUCKET),
            "fileSize": doc.get("fileSize", 0),
            "uploadedBy": doc.get("uploadedBy", ""),
            "username": doc.get("username", ""),
            "totalRows": doc.get("totalRows", 0),
            "createdCount": doc.get("createdCount", 0),
            "updatedCount": doc.get("updatedCount", 0),
            "status": doc.get("status", "success"),
            "message": doc.get("message", ""),
            "resultFile": doc.get("resultFile"),
            "createdAt": doc.get("createdAt", _now()),
        })

    return {
        "items": items,
        "total": total,
        "page": page,
        "pageSize": page_size,
        "totalPages": total_pages,
    }


@router.get("/expiring-contracts/history/{history_id}/download")
async def download_expiring_contracts_history_file(
    history_id: str,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Tải file import ngày đáo hạn gốc đã upload từ S3 MinIO."""
    if not ObjectId.is_valid(history_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mã lịch sử không hợp lệ",
        )

    history = await db[EXPIRING_HISTORY_COLLECTION].find_one({"_id": ObjectId(history_id)})
    if not history:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy lịch sử import ngày đáo hạn",
        )

    s3_key = history.get("s3Key")
    s3_bucket = history.get("s3Bucket", settings.S3_BUCKET)
    filename = history.get("filename", "import_ngay_dao_han.xlsx")

    if not s3_key:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không có thông tin lưu trữ S3 cho file này",
        )

    try:
        s3 = get_s3_client()
        response = s3.get_object(Bucket=s3_bucket, Key=s3_key)

        def iter_chunks():
            for chunk in response["Body"].iter_chunks(chunk_size=1024 * 1024):
                yield chunk

        return StreamingResponse(
            iter_chunks(),
            media_type=response.get("ContentType", "application/octet-stream"),
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
            },
        )
    except Exception as e:
        logger.error(f"Lỗi khi tải file từ S3 cho History ID={history_id}: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể tải file từ S3 MinIO",
        )


@router.get("/expiring-contracts/history/{history_id}/download-result")
async def download_expiring_contracts_result_file(
    history_id: str,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Tải file kết quả ngày đáo hạn kèm cột Thông báo từ S3 MinIO."""
    if not ObjectId.is_valid(history_id):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mã lịch sử không hợp lệ",
        )

    history = await db[EXPIRING_HISTORY_COLLECTION].find_one({"_id": ObjectId(history_id)})
    if not history:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Không tìm thấy lịch sử import ngày đáo hạn",
        )

    result_file = history.get("resultFile")
    if not result_file or not isinstance(result_file, dict) or not result_file.get("s3Key"):
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Phiên này chưa có file kết quả",
        )

    s3_key = result_file.get("s3Key")
    s3_bucket = result_file.get("s3Bucket", settings.S3_BUCKET)
    filename = result_file.get("filename", "ket_qua_ngay_dao_han.xlsx")

    try:
        s3 = get_s3_client()
        response = s3.get_object(Bucket=s3_bucket, Key=s3_key)

        def iter_chunks():
            for chunk in response["Body"].iter_chunks(chunk_size=1024 * 1024):
                yield chunk

        return StreamingResponse(
            iter_chunks(),
            media_type=response.get("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            headers={
                "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
            },
        )
    except Exception as e:
        logger.error(f"Lỗi khi tải file kết quả từ S3 cho History ID={history_id}: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không thể tải file kết quả từ S3 MinIO",
        )


@router.get("/{employee_code}", response_model=ReconciliationRecordResponse)
async def get_reconciliation_record(
    employee_code: str,
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy thông tin đối soát của 1 cộng tác viên."""
    doc = await db[COLLECTION].find_one({"employeeCode": employee_code.strip()})
    if not doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f'Không tìm thấy thông tin đối soát cho CTV mã "{employee_code}"'
        )
    return _to_response(doc)

@router.post("/sync-system-info")
async def sync_bmk_system_info(
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """
    Đồng bộ thông tin số lượng hồ sơ từ collection bmk_ctv_collaborators
    vào trường bmkSystemInfo, createdSource="BMK System" và isBmkSystemExist=True của collection bmk_ctv_reconciliations.
    """
    collab_cursor = db[COLLABORATORS_COLLECTION].find({})
    operations = []
    total_processed = 0
    synced_codes = set()
    now = _now()

    async for collab in collab_cursor:
        code = collab.get("employeeCode") or str(collab.get("_id"))
        if not code:
            continue
        code = str(code).strip()
        synced_codes.add(code)

        system_info = _calculate_system_info(collab)
        full_name = collab.get("fullName", "")
        id_number = collab.get("idNumber") or collab.get("cccd") or None

        operations.append(
            UpdateOne(
                {"employeeCode": code},
                {
                    "$set": {
                        "idNumber": id_number,
                        "bmkSystemInfo": system_info,
                        "createdSource": "BMK System",
                        "isBmkSystemExist": True,
                        "updatedAt": now,
                    },
                    "$setOnInsert": {
                        "_id": code,
                        "employeeCode": code,
                        "fullName": full_name,
                        "departmentLevel1": None,
                        "position": None,
                        "employmentStatus": None,
                        "onboardDate": None,
                        "offboardDate": None,
                        "tpbankInfo": {"contracts": []},
                        "bmkHrInfo": {
                            "contractCount": 0,
                            "idCardCount": 0,
                            "liquidationCount": 0,
                            "taxCommitmentCount": 0,
                        },
                        "reconciliationStatus": "pending",
                        "createdAt": now,
                    },
                },
                upsert=True,
            )
        )

        total_processed += 1

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    # Đánh dấu isBmkSystemExist = False cho các bản ghi trong bmk_ctv_reconciliations không nằm trong bmk_ctv_collaborators
    if synced_codes:
        await db[COLLECTION].update_many(
            {"employeeCode": {"$nin": list(synced_codes)}},
            {"$set": {"isBmkSystemExist": False}}
        )

    return {
        "status": "success",
        "message": f"Đã đồng bộ thành công {total_processed} cộng tác viên từ hệ thống BMK",
        "totalProcessed": total_processed,
    }



@router.post("/import-hr-bmk")
async def import_hr_bmk_data(
    file: UploadFile = File(...),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """
    Import/Update dữ liệu từ file Excel HR BMK vào collection bmk_ctv_reconciliations.
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file Excel (.xlsx, .xlsm)",
        )

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Không đọc được file Excel: {str(e)}",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File Excel không chứa dữ liệu",
        )

    # Tìm dòng tiêu đề chứa 'MNV' hoặc 'MÃ NV'
    header_row_idx = None
    header = []
    for idx, r in enumerate(rows):
        if r and any(cell and ("MNV" in str(cell).upper() or "MÃ NV" in str(cell).upper() or "MÃ NHÂN VIÊN" in str(cell).upper()) for cell in r):
            header_row_idx = idx
            header = [str(c).strip() if c is not None else "" for c in r]
            break

    if header_row_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không tìm thấy dòng tiêu đề chứa cột 'MNV' trong file Excel",
        )

    header_upper = [h.upper() for h in header]

    def find_col_idx(candidates: List[str]) -> Optional[int]:
        for candidate in candidates:
            cand_u = candidate.upper()
            for idx, h in enumerate(header_upper):
                if cand_u in h:
                    return idx
        return None

    code_idx = find_col_idx(["MNV", "MÃ NV", "MÃ NHÂN VIÊN"])
    name_idx = find_col_idx(["HỌ VÀ TÊN", "HỌ TÊN"])
    cccd_idx = find_col_idx(["SỐ CCCD", "CCCD"])
    contract_idx = find_col_idx(["HỢP ĐỒNG", "SL HĐ"])
    
    # Cột CCCD đếm số lượng (thường đứng sau HỢP ĐỒNG)
    id_card_cnt_idx = None
    for idx, h in enumerate(header_upper):
        if "CCCD" in h and idx != cccd_idx:
            id_card_cnt_idx = idx
            break
    if id_card_cnt_idx is None:
        id_card_cnt_idx = cccd_idx

    bbtl_idx = find_col_idx(["BBTL", "BIÊN BẢN THANH LÝ"])
    tax_idx = find_col_idx(["CAM KẾT THUẾ", "CKT"])

    if code_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File thiếu cột 'MNV' (Mã nhân viên)",
        )

    operations = []
    now = _now()

    def parse_int(val) -> int:
        if val is None or val == "":
            return 0
        try:
            return int(float(str(val).strip()))
        except (ValueError, TypeError):
            return 0

    def parse_str(val) -> Optional[str]:
        if val is None:
            return None
        s = str(val).strip()
        if not s or s == "-":
            return None
        return s

    created_count = 0
    updated_count = 0
    total_processed = 0

    # Lấy danh sách các employeeCode đã tồn tại để tránh query 3,000 lần
    existing_codes = set()
    async for doc in db[COLLECTION].find({}, {"employeeCode": 1}):
        code = doc.get("employeeCode")
        if code:
            existing_codes.add(str(code).strip())

    for row in rows[header_row_idx + 1 :]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_code = row[code_idx] if code_idx < len(row) else None
        employee_code = parse_str(raw_code)
        if not employee_code:
            continue

        if employee_code.isdigit() and len(employee_code) < 5:
            employee_code = employee_code.zfill(5)

        full_name = parse_str(row[name_idx] if name_idx is not None and name_idx < len(row) else None) or ""
        id_number = parse_str(row[cccd_idx] if cccd_idx is not None and cccd_idx < len(row) else None)

        contract_cnt = parse_int(row[contract_idx] if contract_idx is not None and contract_idx < len(row) else None)
        id_card_cnt = parse_int(row[id_card_cnt_idx] if id_card_cnt_idx is not None and id_card_cnt_idx < len(row) else None)
        liquidation_cnt = parse_int(row[bbtl_idx] if bbtl_idx is not None and bbtl_idx < len(row) else None)
        tax_cnt = parse_int(row[tax_idx] if tax_idx is not None and tax_idx < len(row) else None)

        bmk_hr_info = {
            "contractCount": contract_cnt,
            "idCardCount": id_card_cnt,
            "liquidationCount": liquidation_cnt,
            "taxCommitmentCount": tax_cnt,
        }

        if employee_code in existing_codes:
            updated_count += 1
        else:
            created_count += 1
            existing_codes.add(employee_code)

        operations.append(
            UpdateOne(
                {"employeeCode": employee_code},
                {
                    "$set": {
                        "idNumber": id_number,
                        "bmkHrInfo": bmk_hr_info,
                        "createdSource": "HR BMK",
                        "updatedAt": now,
                    },
                    "$setOnInsert": {
                        "_id": employee_code,
                        "employeeCode": employee_code,
                        "fullName": full_name,
                        "departmentLevel1": None,
                        "position": None,
                        "employmentStatus": None,
                        "onboardDate": None,
                        "offboardDate": None,
                        "tpbankInfo": {"contracts": []},
                        "bmkSystemInfo": {
                            "contractCount": 0,
                            "idCardCount": 0,
                            "liquidationCount": 0,
                            "taxCommitmentCount": 0,
                        },
                        "reconciliationStatus": "pending",
                        "createdAt": now,
                    },
                },
                upsert=True,
            )
        )

        total_processed += 1

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    return {
        "status": "success",
        "message": f"Đã nhập thành công {total_processed} cộng tác viên từ HR BMK (Tạo mới: {created_count}, Cập nhật: {updated_count})",
        "totalProcessed": total_processed,
        "createdCount": created_count,
        "updatedCount": updated_count,
    }


@router.post("/import-hr-tpbank", response_model=ImportHrTpBankResult)
async def import_hr_tpbank_contracts(
    file: UploadFile = File(...),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """
    Import số lượng Hợp đồng từ file Excel HR TP Bank (template_ctv_hr_tpbank.xlsx).
    Cột mẫu: STT, Mã NV, Họ và tên, Số lượng hợp đồng.
    - Lưu vào TpBankInfo.hrContractCount
    - Nếu CTV chưa tồn tại: tạo mới với employeeCode, fullName, TpBankInfo.hrContractCount, createdSource="HR TP Bank"
    - Nếu CTV đã tồn tại: cập nhật TpBankInfo.hrContractCount
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file Excel (.xlsx, .xlsm)",
        )

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Không đọc được file Excel: {str(e)}",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File Excel không chứa dữ liệu",
        )

    # Tìm dòng tiêu đề chứa 'MNV', 'MÃ NV', hoặc 'MÃ NHÂN VIÊN'
    header_row_idx = None
    header = []
    for idx, r in enumerate(rows):
        if r and any(cell and ("MNV" in str(cell).upper() or "MÃ NV" in str(cell).upper() or "MÃ NHÂN VIÊN" in str(cell).upper()) for cell in r):
            header_row_idx = idx
            header = [str(c).strip() if c is not None else "" for c in r]
            break

    if header_row_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không tìm thấy dòng tiêu đề chứa cột 'Mã NV' trong file Excel",
        )

    header_upper = [h.upper() for h in header]

    def find_col_idx(candidates: List[str]) -> Optional[int]:
        for candidate in candidates:
            cand_u = candidate.upper()
            for idx, h in enumerate(header_upper):
                if cand_u in h:
                    return idx
        return None

    code_idx = find_col_idx(["MÃ NV", "MNV", "MÃ NHÂN VIÊN"])
    name_idx = find_col_idx(["HỌ VÀ TÊN", "HỌ TÊN", "TÊN"])
    contract_idx = find_col_idx(["SỐ LƯỢNG HỢP ĐỒNG", "SL HỢP ĐỒNG", "SL HĐ", "HỢP ĐỒNG"])

    if code_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File thiếu cột 'Mã NV'",
        )

    def parse_int(val) -> Optional[int]:
        if val is None or str(val).strip() in ["", "-", "None", "null"]:
            return None
        try:
            return int(float(str(val).strip()))
        except (ValueError, TypeError):
            return None

    def parse_str(val) -> Optional[str]:
        if val is None:
            return None
        s = str(val).strip()
        if not s or s == "-":
            return None
        return s

    operations = []
    now = _now()
    created_count = 0
    updated_count = 0
    total_processed = 0

    existing_docs = {}
    async for doc in db[COLLECTION].find({}, {
        "employeeCode": 1,
        "fullName": 1,
        "tpbankInfo": 1,
        "bmkHrInfo": 1,
        "bmkSystemInfo": 1,
        "employmentStatus": 1,
        "isSynced": 1,
        "result": 1,
    }):
        code = doc.get("employeeCode")
        if code:
            existing_docs[str(code).strip()] = doc

    for row in rows[header_row_idx + 1:]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_code = row[code_idx] if code_idx < len(row) else None
        employee_code = parse_str(raw_code)
        if not employee_code:
            continue

        if employee_code.isdigit() and len(employee_code) < 5:
            employee_code = employee_code.zfill(5)

        full_name = parse_str(row[name_idx] if name_idx is not None and name_idx < len(row) else None) or ""
        contract_cnt = parse_int(row[contract_idx] if contract_idx is not None and contract_idx < len(row) else None)

        total_processed += 1

        if employee_code in existing_docs:
            updated_count += 1
            ex_doc = existing_docs[employee_code]
            update_fields = {
                "tpbankInfo.hrContractCount": contract_cnt,
                "updatedAt": now,
            }
            if ex_doc.get("isSynced") or ex_doc.get("result"):
                est_cnt = (ex_doc.get("tpbankInfo") or {}).get("estimatedContractCount")
                bmk_hr = ex_doc.get("bmkHrInfo") or {}
                bmk_sys = ex_doc.get("bmkSystemInfo") or {}
                emp_status = ex_doc.get("employmentStatus")
                result_obj = _calculate_reconciliation_result(
                    estimated_contract_count=est_cnt,
                    hr_contract_count=contract_cnt,
                    bmk_hr_info=bmk_hr,
                    bmk_system_info=bmk_sys,
                    employment_status=emp_status,
                )
                update_fields["result"] = result_obj
            if full_name:
                operations.append(
                    UpdateOne(
                        {"employeeCode": employee_code, "$or": [{"fullName": ""}, {"fullName": None}, {"fullName": {"$exists": False}}]},
                        {"$set": {"fullName": full_name}}
                    )
                )
            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {"$set": update_fields}
                )
            )
        else:
            created_count += 1
            existing_codes.add(employee_code)
            new_doc = {
                "_id": employee_code,
                "employeeCode": employee_code,
                "fullName": full_name,
                "createdSource": "HR TP Bank",
                "departmentLevel1": None,
                "position": None,
                "employmentStatus": None,
                "onboardDate": None,
                "offboardDate": None,
                "tpbankInfo": {
                    "estimatedContractCount": None,
                    "hrContractCount": contract_cnt,
                    "contracts": [],
                },
                "bmkHrInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "bmkSystemInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "reconciliationStatus": "pending",
                "createdAt": now,
                "updatedAt": now,
            }
            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {"$setOnInsert": new_doc},
                    upsert=True,
                )
            )

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    return {
        "status": "success",
        "message": f"Đã nhập thành công số lượng HĐ TP Bank: {total_processed} CTV (Tạo mới: {created_count}, Cập nhật: {updated_count})",
        "totalProcessed": total_processed,
        "createdCount": created_count,
        "updatedCount": updated_count,
    }


def _calculate_tpbank_contract_count(onboard_dt: Optional[datetime], offboard_dt: Optional[datetime]) -> Optional[int]:
    if not onboard_dt:
        return None
    end_dt = offboard_dt if offboard_dt else datetime.now()
    if end_dt < onboard_dt:
        return 0

    full_months = (end_dt.year - onboard_dt.year) * 12 + (end_dt.month - onboard_dt.month)
    # Quy tắc: nếu ngày kết thúc chưa vượt quá ngày vào (<= onboard_dt.day),
    # tức vừa chạm đúng mốc tròn tháng, thì chưa phát sinh sang hợp đồng tiếp theo
    if end_dt.day <= onboard_dt.day:
        full_months -= 1

    if full_months < 0:
        full_months = 0

    return (full_months // 6) + 1


def _calculate_reconciliation_result(
    estimated_contract_count: Optional[int],
    hr_contract_count: Optional[int],
    bmk_hr_info: dict,
    bmk_system_info: dict,
    employment_status: Optional[str],
) -> dict:
    sys_contract = (bmk_system_info or {}).get("contractCount", 0) or 0
    hr_contract = (bmk_hr_info or {}).get("contractCount", 0) or 0
    total_bmk_contract = sys_contract + hr_contract

    # 1. Đầu tiên kiểm tra nếu TpBankInfo.estimatedContractCount != TpBankInfo.hrContractCount -> cảnh báo Lệch Bank
    if estimated_contract_count != hr_contract_count:
        contract_res = "warn"
    else:
        # 2. Nếu TpBankInfo.estimatedContractCount == TpBankInfo.hrContractCount thì thực hiện tiếp logic hiện tại
        if estimated_contract_count is not None and estimated_contract_count <= total_bmk_contract:
            contract_res = "success"
        else:
            contract_res = "failed"

    # Đối soát idCard: nếu bmkSystemInfo.idCardCount = 1 hoặc bmkHrInfo.idCardCount = 1 -> success còn lại failed
    sys_idcard = (bmk_system_info or {}).get("idCardCount", 0) or 0
    hr_idcard = (bmk_hr_info or {}).get("idCardCount", 0) or 0
    if sys_idcard == 1 or hr_idcard == 1:
        idcard_res = "success"
    else:
        idcard_res = "failed"

    # Đối soát liquidation:
    # 1. Kết quả "Khớp" (success):
    #    - Nếu Nhân viên Nghỉ việc: Bắt buộc phải có BBTL trên Hệ thống BMK (bmkSystemInfo.liquidationCount == 1) hoặc trong file HR BMK (bmkHrInfo.liquidationCount == 1)
    #    - Nếu Nhân viên đang hiệu lực: Không có BBTL trên Hệ thống BMK (bmkSystemInfo.liquidationCount == 0) và trong file HR BMK (bmkHrInfo.liquidationCount == 0)
    # 2. Kết quả "Lệch" (failed): Các trường hợp còn lại
    sys_liq = (bmk_system_info or {}).get("liquidationCount", 0) or 0
    hr_liq = (bmk_hr_info or {}).get("liquidationCount", 0) or 0
    if employment_status == "Nghỉ việc":
        if sys_liq == 1 or hr_liq == 1:
            liq_res = "success"
        else:
            liq_res = "failed"
    else:
        if sys_liq == 0 and hr_liq == 0:
            liq_res = "success"
        else:
            liq_res = "failed"

    sys_tax = (bmk_system_info or {}).get("taxCommitmentCount", 0) or 0
    hr_tax = (bmk_hr_info or {}).get("taxCommitmentCount", 0) or 0
    tax_res = "success" if (sys_tax == 1 or hr_tax == 1) else "failed"

    return {
        "contract": contract_res,
        "idCard": idcard_res,
        "liquidation": liq_res,
        "taxCommitment": tax_res,
    }


@router.post("/reconcile-tpbank")
async def reconcile_tpbank_data(
    file: UploadFile = File(...),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """
    Upload file Đối soát TP Bank, thực hiện đối soát và cập nhật collection bmk_ctv_reconciliations.
    1. Tại thời điểm upload update tất cả row có isSynced = False
    2. Đọc MNV, HỌ VÀ TÊN, NGÀY GIA NHẬP, NGÀY THÔI VIỆC
    3. Nếu chưa có -> tạo mới với createdSource="TP Bank", isSynced=True
    4. Nếu đã có -> update NGÀY GIA NHẬP, NGÀY THÔI VIỆC, isSynced=True
    5. Tính số HĐ dự tính lưu vào tpbankInfo.estimatedContractCount
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file Excel (.xlsx, .xlsm)",
        )

    # 1. Update tất cả bản ghi hiện tại thành isSynced = False và result = None
    await db[COLLECTION].update_many({}, {"$set": {"isSynced": False, "result": None}})

    content = await file.read()

    # Upload file đối soát gốc lên S3 MinIO
    now_ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    s3_key = f"reconcile_tpbank/{now_ts}_{file.filename}"
    s3_bucket = settings.S3_BUCKET
    try:
        upload_to_s3(content, s3_key)
    except Exception as e:
        logger.warning(f"Không thể lưu file lên S3 MinIO: {str(e)}")

    try:
        wb = load_workbook(BytesIO(content))
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Không đọc được file Excel: {str(e)}",
        )

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 2:
        wb.close()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File Excel không chứa dữ liệu",
        )

    header_row_idx = None
    header = []
    for idx, r in enumerate(rows):
        if r and any(cell and ("MNV" in str(cell).upper() or "MÃ NV" in str(cell).upper() or "MÃ NHÂN VIÊN" in str(cell).upper()) for cell in r):
            header_row_idx = idx
            header = [str(c).strip() if c is not None else "" for c in r]
            break

    if header_row_idx is None:
        wb.close()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không tìm thấy dòng tiêu đề chứa cột 'MNV' trong file Excel",
        )

    header_upper = [h.upper() for h in header]

    def find_col_idx(candidates: List[str]) -> Optional[int]:
        for candidate in candidates:
            cand_u = candidate.upper()
            for idx, h in enumerate(header_upper):
                if cand_u in h:
                    return idx
        return None

    code_idx = find_col_idx(["MNV", "MÃ NV", "MÃ NHÂN VIÊN"])
    name_idx = find_col_idx(["HỌ VÀ TÊN", "HỌ TÊN"])
    onboard_idx = find_col_idx(["NGÀY GIA NHẬP", "NGÀY VÀO"])
    offboard_idx = find_col_idx(["NGÀY THÔI VIỆC", "NGÀY NGHỈ"])

    if code_idx is None:
        wb.close()
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File thiếu cột 'MNV' (Mã nhân viên)",
        )

    # Khởi tạo tiêu đề 11 cột kết quả đối soát mới (bắt đầu từ cột 11 / K)
    new_headers = [
        (11, "HĐ dự tính"),
        (12, "HĐ từ bank"),
        (13, "HĐ HR BMK"),
        (14, "CCCD HR BMK"),
        (15, "BBTL HR BMK"),
        (16, "HĐ System BMK"),
        (17, "CCCD System BMK"),
        (18, "BBTL System BMK"),
        (19, "Đối soát HĐ"),
        (20, "Đối soát CCCD"),
        (21, "Đối soát BBTL"),
    ]
    header_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    header_fill_tpbank = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
    header_fill_hr = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    header_fill_sys = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    header_fill_res = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

    thin_side = Side(style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    data_align_center = Alignment(horizontal="center", vertical="center")
    font_not_exist = Font(name="Calibri", size=10, italic=True, color="DC2626")
    font_match = Font(name="Calibri", size=10, bold=True, color="166534")
    font_warn = Font(name="Calibri", size=10, bold=True, color="B45309")
    font_mismatch = Font(name="Calibri", size=10, bold=True, color="DC2626")

    header_row_num = header_row_idx + 1
    for col_idx_new, title in new_headers:
        c = ws.cell(row=header_row_num, column=col_idx_new, value=title)
        c.font = header_font
        c.alignment = header_align
        c.border = thin_border
        if col_idx_new in [11, 12]:
            c.fill = header_fill_tpbank
        elif col_idx_new in [13, 14, 15]:
            c.fill = header_fill_hr
        elif col_idx_new in [16, 17, 18]:
            c.fill = header_fill_sys
        else:
            c.fill = header_fill_res

    existing_docs = {}
    async for doc in db[COLLECTION].find({}, {
        "employeeCode": 1,
        "fullName": 1,
        "isBmkSystemExist": 1,
        "tpbankInfo": 1,
        "bmkHrInfo": 1,
        "bmkSystemInfo": 1,
    }):
        code = doc.get("employeeCode")
        if code:
            existing_docs[str(code).strip()] = doc

    operations = []
    now = _now()
    created_count = 0
    updated_count = 0
    total_processed = 0

    for r_offset, row in enumerate(rows[header_row_idx + 1:]):
        excel_row_num = header_row_idx + 2 + r_offset

        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_code = row[code_idx] if code_idx < len(row) else None
        if raw_code is None:
            continue
        employee_code = str(raw_code).strip()
        if not employee_code or employee_code.upper() in ["MNV", "MÃ NV"]:
            continue

        if employee_code.isdigit() and len(employee_code) < 5:
            employee_code = employee_code.zfill(5)

        full_name = str(row[name_idx]).strip() if name_idx is not None and name_idx < len(row) and row[name_idx] is not None else ""
        raw_onboard = row[onboard_idx] if onboard_idx is not None and onboard_idx < len(row) else None
        raw_offboard = row[offboard_idx] if offboard_idx is not None and offboard_idx < len(row) else None

        onboard_str, onboard_dt = _parse_excel_date(raw_onboard)
        offboard_str, offboard_dt = _parse_excel_date(raw_offboard)

        # quy tắc: Nếu ngày gia nhập trước ngày 1/4/2025 (dd/MM/yyyy) thì Ngày gia nhập = 1/4/2025
        if onboard_dt and onboard_dt < datetime(2025, 4, 1):
            onboard_dt = datetime(2025, 4, 1)
            onboard_str = "2025-04-01"

        # quy tắc: nếu nhân viên có giá trị ở cột NGÀY THÔI VIỆC là null hoặc N/a hoặc ngày 1/0/1900 (4 số cuối 1900)
        # thì ngày thôi việc là null và đang làm việc employmentStatus = "Hiệu lực" còn lại "Nghỉ việc"
        if offboard_dt is None:
            offboard_str = None
            employment_status = "Hiệu lực"
        else:
            employment_status = "Nghỉ việc"

        contract_count = _calculate_tpbank_contract_count(onboard_dt, offboard_dt)

        if employee_code in existing_docs:
            updated_count += 1
            existing_doc = existing_docs[employee_code]
            tpbank_info = existing_doc.get("tpbankInfo") or {}
            contracts = tpbank_info.get("contracts") or []
            bmk_hr_info = existing_doc.get("bmkHrInfo") or {}
            bmk_system_info = existing_doc.get("bmkSystemInfo") or {}
            is_sys_exist = existing_doc.get("isBmkSystemExist", False)
            hr_bank_cnt = tpbank_info.get("hrContractCount")

            result_obj = _calculate_reconciliation_result(
                estimated_contract_count=contract_count,
                hr_contract_count=hr_bank_cnt,
                bmk_hr_info=bmk_hr_info,
                bmk_system_info=bmk_system_info,
                employment_status=employment_status,
            )

            new_tpbank_info = {
                "estimatedContractCount": contract_count,
                "hrContractCount": hr_bank_cnt,
                "contracts": contracts,
            }
            update_set = {
                "onboardDate": onboard_str,
                "offboardDate": offboard_str,
                "employmentStatus": employment_status,
                "isSynced": True,
                "tpbankInfo": new_tpbank_info,
                "result": result_obj,
                "updatedAt": now,
            }
            if full_name and not existing_doc.get("fullName"):
                update_set["fullName"] = full_name

            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {"$set": update_set}
                )
            )
        else:
            created_count += 1
            existing_docs[employee_code] = {"employeeCode": employee_code, "fullName": full_name, "isBmkSystemExist": False}
            new_tpbank_info = {
                "estimatedContractCount": contract_count,
                "hrContractCount": None,
                "contracts": [],
            }
            empty_info = {
                "contractCount": 0,
                "idCardCount": 0,
                "liquidationCount": 0,
                "taxCommitmentCount": 0,
            }
            bmk_hr_info = empty_info
            bmk_system_info = empty_info
            is_sys_exist = False
            hr_bank_cnt = None

            result_obj = _calculate_reconciliation_result(
                estimated_contract_count=contract_count,
                hr_contract_count=None,
                bmk_hr_info=empty_info,
                bmk_system_info=empty_info,
                employment_status=employment_status,
            )
            operations.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {
                        "$set": {
                            "onboardDate": onboard_str,
                            "offboardDate": offboard_str,
                            "employmentStatus": employment_status,
                            "isSynced": True,
                            "tpbankInfo": new_tpbank_info,
                            "result": result_obj,
                            "updatedAt": now,
                        },
                        "$setOnInsert": {
                            "_id": employee_code,
                            "employeeCode": employee_code,
                            "fullName": full_name,
                            "idNumber": None,
                            "createdSource": "TP Bank",
                            "isBmkSystemExist": False,
                            "departmentLevel1": None,
                            "position": None,
                            "bmkHrInfo": empty_info,
                            "bmkSystemInfo": empty_info,
                            "reconciliationStatus": "pending",
                            "createdAt": now,
                        },
                    },
                    upsert=True,
                )
            )

        # Điền các cột kết quả đối soát vào Excel cho row hiện tại
        hr_cnt_val = (bmk_hr_info or {}).get("contractCount", 0) or 0
        hr_idcard_val = (bmk_hr_info or {}).get("idCardCount", 0) or 0
        hr_liq_val = (bmk_hr_info or {}).get("liquidationCount", 0) or 0

        sys_cnt_val = (bmk_system_info or {}).get("contractCount", 0) or 0 if is_sys_exist else "Nhân viên không tồn tại"
        sys_idcard_val = (bmk_system_info or {}).get("idCardCount", 0) or 0 if is_sys_exist else "Nhân viên không tồn tại"
        sys_liq_val = (bmk_system_info or {}).get("liquidationCount", 0) or 0 if is_sys_exist else "Nhân viên không tồn tại"

        contract_res = result_obj.get("contract")
        idcard_res = result_obj.get("idCard")
        liq_res = result_obj.get("liquidation")

        contract_text = "Khớp" if contract_res == "success" else ("Lệch Bank" if contract_res == "warn" else "Lệch")
        contract_font = font_warn if contract_res == "warn" else (font_match if contract_res == "success" else font_mismatch)

        idcard_text = "Khớp" if idcard_res == "success" else "Lệch"
        idcard_font = font_match if idcard_res == "success" else font_mismatch

        liq_text = "Khớp" if liq_res == "success" else "Lệch"
        liq_font = font_match if liq_res == "success" else font_mismatch

        row_cells = [
            (11, contract_count, data_font),
            (12, hr_bank_cnt if hr_bank_cnt is not None else "", data_font),
            (13, hr_cnt_val, data_font),
            (14, hr_idcard_val, data_font),
            (15, hr_liq_val, data_font),
            (16, sys_cnt_val, data_font if is_sys_exist else font_not_exist),
            (17, sys_idcard_val, data_font if is_sys_exist else font_not_exist),
            (18, sys_liq_val, data_font if is_sys_exist else font_not_exist),
            (19, contract_text, contract_font),
            (20, idcard_text, idcard_font),
            (21, liq_text, liq_font),
        ]
        for col_num, val, c_font in row_cells:
            cell = ws.cell(row=excel_row_num, column=col_num, value=val)
            cell.font = c_font
            cell.alignment = data_align_center
            cell.border = thin_border

        total_processed += 1

        if len(operations) >= 500:
            await db[COLLECTION].bulk_write(operations, ordered=False)
            operations = []

    if operations:
        await db[COLLECTION].bulk_write(operations, ordered=False)

    # Định dạng độ rộng các cột mới
    col_widths = {
        11: 14, 12: 14,
        13: 14, 14: 15, 15: 15,
        16: 22, 17: 22, 18: 22,
        19: 15, 20: 15, 21: 15,
    }
    for col_num, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_num)].width = width

    # Lưu workbook kết quả đối soát vào buffer và upload lên S3 MinIO
    result_file_info = None
    try:
        output_stream = BytesIO()
        wb.save(output_stream)
        wb.close()
        result_content = output_stream.getvalue()

        clean_name = re.sub(r"[^\w\.-]", "_", file.filename)
        result_filename = f"ket_qua_doi_soat_{clean_name}"
        result_s3_key = f"reconcile_results/{now_ts}_{clean_name}"
        upload_to_s3(result_content, result_s3_key)
        result_file_info = {
            "filename": result_filename,
            "s3Key": result_s3_key,
            "s3Bucket": s3_bucket,
            "fileSize": len(result_content),
        }
    except Exception as e:
        logger.error(f"Không thể lưu file kết quả đối soát lên S3 MinIO: {str(e)}")

    non_empty_rows = [
        r for r in rows[header_row_idx + 1:]
        if r and any(cell is not None and str(cell).strip() != "" for cell in r)
    ]
    total_rows_in_file = len(non_empty_rows)
    success_rows = total_processed
    failed_rows = max(0, total_rows_in_file - success_rows)

    # Thống kê kết quả đối soát tổng hợp của phiên này
    total_success = await db[COLLECTION].count_documents({
        "isSynced": True,
        "result.contract": "success",
        "result.idCard": "success",
        "result.liquidation": "success",
    })
    total_warn_bank = await db[COLLECTION].count_documents({
        "isSynced": True,
        "result.contract": "warn",
    })
    total_mismatch_contract = await db[COLLECTION].count_documents({
        "isSynced": True,
        "result.contract": "failed",
    })
    total_mismatch_idcard = await db[COLLECTION].count_documents({
        "isSynced": True,
        "result.idCard": "failed",
    })
    total_mismatch_liquidation = await db[COLLECTION].count_documents({
        "isSynced": True,
        "result.liquidation": "failed",
    })

    # Lưu lịch sử đối soát vào collection bmk_ctv_reconciliation_history
    uploaded_by = current_user.get("name") or current_user.get("username", "Unknown")
    username = current_user.get("username", "Unknown")
    history_doc = {
        "filename": file.filename,
        "s3Key": s3_key,
        "s3Bucket": s3_bucket,
        "fileSize": len(content),
        "uploadedBy": uploaded_by,
        "username": username,
        "totalRows": total_rows_in_file,
        "successRows": success_rows,
        "failedRows": failed_rows,
        "status": "success",
        "message": f"Đối soát thành công {total_processed} bản ghi từ file TP Bank (Tạo mới: {created_count}, Cập nhật: {updated_count})",
        "stats": {
            "totalSuccess": total_success,
            "totalWarnBank": total_warn_bank,
            "totalMismatchContract": total_mismatch_contract,
            "totalMismatchIdCard": total_mismatch_idcard,
            "totalMismatchLiquidation": total_mismatch_liquidation,
        },
        "resultFile": result_file_info,
        "createdAt": now,
    }
    await db[HISTORY_COLLECTION].insert_one(history_doc)

    return {
        "status": "success",
        "message": f"Đối soát thành công {total_processed} bản ghi từ file TP Bank (Tạo mới: {created_count}, Cập nhật: {updated_count})",
        "totalProcessed": total_processed,
        "createdCount": created_count,
        "updatedCount": updated_count,
    }




