import re
from datetime import datetime, date, timezone
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import List, Optional

from dateutil.relativedelta import relativedelta
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, status
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pymongo import UpdateOne

from app.core.database import get_db
from app.core.logging import get_logger
from app.core.security import get_current_user
from app.models.expense import (
    ExpenseStatsResponse,
    ExpenseImportResponse,
    MonthlyCollaboratorCountStat,
    MonthlyPaymentAmountStat,
)

logger = get_logger(__name__)

router = APIRouter(prefix="/api/expenses", tags=["Collaborator Expenses"])

COLLABORATOR_EXPENSES_COLLECTION = "bmk_ctv_collaborator_expenses"
EXPENSE_DETAILS_COLLECTION = "bmk_ctv_expense_details"
RECONCILIATION_COLLECTION = "bmk_ctv_reconciliations"

TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / "import_thanh_toan_phi_CTV.xlsx"
EXPORT_TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "templates" / "export_thanh_toan_phi_CTV.xlsx"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _actor_name(current_user: dict) -> str:
    return current_user.get("name") or current_user.get("fullName") or current_user.get("username", "System")


def parse_employee_code(val) -> Optional[str]:
    if val is None:
        return None
    s = str(val).strip()
    if not s or s.lower() in ("-", "none", "null"):
        return None
    if s.endswith(".0"):
        s = s[:-2]
    if s.isdigit() and len(s) < 5:
        s = s.zfill(5)
    return s


def parse_str_val(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ("none", "null", "-"):
        return ""
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def parse_float_val(val) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace(" ", "").replace("₫", "").replace("VND", "").replace("vnd", "")
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def parse_date_str(val) -> Optional[str]:
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s.lower() in ("-", "none", "null"):
        return None
    if " " in s:
        s = s.split(" ")[0].strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return None


def calculate_period(payment_date_str: str) -> str:
    """Calculate period 'YYYYMM' from payment date 'YYYY-MM-DD'."""
    try:
        parts = payment_date_str.split("-")
        if len(parts) >= 2:
            return f"{parts[0]}{parts[1]}"
    except Exception:
        pass
    return "UNKNOWN"


def format_period_label(period_str: str) -> str:
    """Format 'YYYYMM' to 'MM/YYYY'."""
    if len(period_str) == 6 and period_str.isdigit():
        return f"{period_str[4:]}/{period_str[:4]}"
    return period_str


def format_display_date(date_str: Optional[str]) -> str:
    """Format 'YYYY-MM-DD' to 'DD/MM/YYYY'."""
    if not date_str:
        return ""
    s = str(date_str).strip()
    if not s or s.lower() in ("none", "null", "-"):
        return ""
    if " " in s:
        s = s.split(" ")[0].strip()
    try:
        parts = s.split("-")
        if len(parts) == 3 and len(parts[0]) == 4:
            return f"{parts[2]}/{parts[1]}/{parts[0]}"
    except Exception:
        pass
    return s


def add_six_months(date_str: Optional[str]) -> Optional[str]:
    if not date_str:
        return None
    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d").date()
        exp_dt = dt + relativedelta(months=6)
        return exp_dt.strftime("%Y-%m-%d")
    except Exception:
        return None


@router.get("/template")
async def download_template(current_user: dict = Depends(get_current_user)):
    """Tải file mẫu Excel import thanh toán phí CTV."""
    if not TEMPLATE_PATH.exists():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="File mẫu import_thanh_toan_phi_CTV.xlsx không tồn tại trên hệ thống",
        )
    return FileResponse(
        path=str(TEMPLATE_PATH),
        filename="import_thanh_toan_phi_CTV.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/periods")
async def list_periods(
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách các kỳ thanh toán duy nhất (sorted mới nhất -> cũ nhất)."""
    periods = await db[EXPENSE_DETAILS_COLLECTION].distinct("period")
    periods = [p for p in periods if p and p != "UNKNOWN"]
    periods.sort(reverse=True)
    return {"periods": periods}


@router.get("/locations")
async def list_locations(
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách các nơi làm việc duy nhất."""
    locations = await db[EXPENSE_DETAILS_COLLECTION].distinct("workLocation")
    locations = [loc for loc in locations if loc and loc.strip()]
    locations.sort()
    return {"locations": locations}


@router.get("/stats", response_model=ExpenseStatsResponse)
async def get_expense_stats(
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy thống kê tổng quan và dữ liệu cho 2 biểu đồ:
    1. Số lượng CTV được thanh toán theo từng tháng
    2. Số tiền BMK thanh toán theo từng tháng
    """
    total_collabs = await db[COLLABORATOR_EXPENSES_COLLECTION].count_documents({})
    total_tx = await db[EXPENSE_DETAILS_COLLECTION].count_documents({})

    pipeline_summary = [
        {
            "$group": {
                "_id": None,
                "totalPaidAmount": {"$sum": "$netIncome"},
                "totalTaxableAmount": {"$sum": "$taxableIncome"},
                "totalTaxAmount": {"$sum": "$personalIncomeTax"},
            }
        }
    ]
    summary_res = await db[EXPENSE_DETAILS_COLLECTION].aggregate(pipeline_summary).to_list(1)
    total_paid = summary_res[0]["totalPaidAmount"] if summary_res else 0.0
    total_taxable = summary_res[0]["totalTaxableAmount"] if summary_res else 0.0
    total_tax = summary_res[0]["totalTaxAmount"] if summary_res else 0.0

    # Biểu đồ 1: Số CTV duy nhất được thanh toán theo từng tháng
    pipeline_monthly_collabs = [
        {
            "$group": {
                "_id": {"period": "$period", "employeeCode": "$employeeCode"}
            }
        },
        {
            "$group": {
                "_id": "$_id.period",
                "count": {"$sum": 1}
            }
        },
        {"$sort": {"_id": 1}}
    ]
    monthly_collabs_res = await db[EXPENSE_DETAILS_COLLECTION].aggregate(pipeline_monthly_collabs).to_list(100)
    collab_counts_by_month: List[MonthlyCollaboratorCountStat] = []
    for item in monthly_collabs_res:
        p = item["_id"]
        if p and p != "UNKNOWN":
            collab_counts_by_month.append(
                MonthlyCollaboratorCountStat(
                    period=p,
                    label=format_period_label(p),
                    count=item["count"],
                )
            )

    # Biểu đồ 2: Số tiền thanh toán theo từng tháng (Net income, Taxable, Tax)
    pipeline_monthly_amount = [
        {
            "$group": {
                "_id": "$period",
                "totalNetIncome": {"$sum": "$netIncome"},
                "totalTaxableIncome": {"$sum": "$taxableIncome"},
                "totalPersonalIncomeTax": {"$sum": "$personalIncomeTax"},
            }
        },
        {"$sort": {"_id": 1}}
    ]
    monthly_amount_res = await db[EXPENSE_DETAILS_COLLECTION].aggregate(pipeline_monthly_amount).to_list(100)
    payment_amounts_by_month: List[MonthlyPaymentAmountStat] = []
    for item in monthly_amount_res:
        p = item["_id"]
        if p and p != "UNKNOWN":
            payment_amounts_by_month.append(
                MonthlyPaymentAmountStat(
                    period=p,
                    label=format_period_label(p),
                    totalNetIncome=item.get("totalNetIncome", 0.0),
                    totalTaxableIncome=item.get("totalTaxableIncome", 0.0),
                    totalPersonalIncomeTax=item.get("totalPersonalIncomeTax", 0.0),
                )
            )

    latest_period = None
    if collab_counts_by_month:
        latest_period = collab_counts_by_month[-1].period

    return ExpenseStatsResponse(
        totalPaidAmount=round(total_paid, 2),
        totalTaxableAmount=round(total_taxable, 2),
        totalTaxAmount=round(total_tax, 2),
        totalPaymentTransactions=total_tx,
        totalCollaborators=total_collabs,
        latestPeriod=latest_period,
        collaboratorCountsByMonth=collab_counts_by_month,
        paymentAmountsByMonth=payment_amounts_by_month,
    )


@router.get("/details")
async def list_expense_details(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: Optional[str] = Query(None),
    period: Optional[str] = Query(None),
    work_location: Optional[str] = Query(None),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách chi tiết chi phí thanh toán cho CTV."""
    query = {}
    if period and period.strip():
        query["period"] = period.strip()
    if work_location and work_location.strip():
        query["workLocation"] = work_location.strip()
    if search and search.strip():
        s = search.strip()
        reg = {"$regex": re.escape(s), "$options": "i"}
        query["$or"] = [
            {"employeeCode": reg},
            {"fullName": reg},
            {"taxCode": reg},
            {"bankAccount.accountNumber": reg},
            {"workLocation": reg},
        ]

    total = await db[EXPENSE_DETAILS_COLLECTION].count_documents(query)
    skip = (page - 1) * page_size
    cursor = (
        db[EXPENSE_DETAILS_COLLECTION]
        .find(query)
        .sort([("paymentDate", -1), ("employeeCode", 1)])
        .skip(skip)
        .limit(page_size)
    )

    items = []
    async for doc in cursor:
        doc["id"] = str(doc.get("_id"))
        doc.pop("_id", None)
        items.append(doc)

    total_pages = max(1, (total + page_size - 1) // page_size)
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
    }


@router.get("/collaborators")
async def list_expense_collaborators(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: Optional[str] = Query(None),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Lấy danh sách tổng hợp CTV chi phí (collection chính)."""
    query = {}
    if search and search.strip():
        s = search.strip()
        reg = {"$regex": re.escape(s), "$options": "i"}
        query["$or"] = [
            {"employeeCode": reg},
            {"fullName": reg},
            {"taxCode": reg},
        ]

    total = await db[COLLABORATOR_EXPENSES_COLLECTION].count_documents(query)
    skip = (page - 1) * page_size
    cursor = (
        db[COLLABORATOR_EXPENSES_COLLECTION]
        .find(query)
        .sort([("totalPaymentCount", -1), ("employeeCode", 1)])
        .skip(skip)
        .limit(page_size)
    )

    items = []
    async for doc in cursor:
        doc["id"] = str(doc.get("_id"))
        doc.pop("_id", None)
        items.append(doc)

    total_pages = max(1, (total + page_size - 1) // page_size)
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "total_pages": total_pages,
    }


@router.post("/import", response_model=ExpenseImportResponse)
async def import_expenses(
    file: UploadFile = File(...),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Import dữ liệu thanh toán phí CTV từ file Excel:
    1. Kiểm tra tồn tại ở bmk_ctv_reconciliations: Nếu chưa có thì tạo mới 1 record cơ bản.
    2. Upsert dữ liệu vào collection detail (bmk_ctv_expense_details) với unique key: employeeCode + paymentDate.
    3. Cập nhật thông tin và số lần nhận thanh toán vào collection chính (bmk_ctv_collaborator_expenses).
    """
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file Excel (.xlsx, .xlsm)",
        )

    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Không đọc được file Excel: {str(e)}",
        )

    sheet_name = "Gui BMK" if "Gui BMK" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File Excel không chứa dữ liệu",
        )

    # Tìm dòng tiêu đề chính
    header_row_idx = None
    header = []
    for idx, r in enumerate(rows):
        if r and any(
            cell and any(k in str(cell).upper() for k in ["MNV", "MÃ NV", "MÃ NHÂN VIÊN", "MÃ CTV", "EMPLOYEE CODE"])
            for cell in r
        ):
            header_row_idx = idx
            header = [str(c).strip() if c is not None else "" for c in r]
            break

    if header_row_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không tìm thấy dòng tiêu đề chứa cột 'MNV' trong file Excel",
        )

    header_upper = [h.upper() for h in header]

    def find_col(candidates: List[str]) -> Optional[int]:
        for candidate in candidates:
            cand_u = candidate.upper()
            for idx, h in enumerate(header_upper):
                if cand_u in h:
                    return idx
        return None

    code_idx = find_col(["MNV", "MÃ NV", "MÃ NHÂN VIÊN", "MÃ CTV", "MA NV"])
    name_idx = find_col(["HỌ VÀ TÊN", "HỌ TÊN", "HO VA TEN", "FULL NAME", "NAME"])
    tax_idx = find_col(["MÃ SỐ THUẾ", "MA SO THUE", "MST", "TAX CODE"])
    acc_name_idx = find_col(["TÊN TÀI KHOẢN", "TEN TAI KHOAN", "TÊN TK", "ACCOUNT NAME"])
    acc_num_idx = find_col(["SỐ TÀI KHOẢN", "SO TAI KHOAN", "SỐ TK", "ACCOUNT NUMBER"])
    bank_idx = find_col(["NGÂN HÀNG", "NGAN HANG", "BANK"])
    taxable_idx = find_col(["TỔNG TN", "TONG TN", "THU NHẬP CHỊU THUẾ", "TN CHỊU THUẾ"])
    tax_tncn_idx = find_col(["THUẾ TNCN", "THUE TNCN", "TNCN"])
    net_idx = find_col(["THỰC LĨNH", "THUC LINH", "NET"])
    work_loc_idx = find_col(["NƠI LÀM VIỆC", "NOI LAM VIEC", "LOCATION"])
    pay_date_idx = find_col(["NGÀY BMK TT", "NGAY BMK TT", "NGÀY TT", "NGAY TT", "NGÀY THANH TOÁN"])
    pay_content_idx = find_col(["NỘI DUNG TT", "NOI DUNG TT", "NỘI DUNG", "NOI DUNG"])
    onboard_idx = find_col(["NGÀY GIA NHẬP", "NGAY GIA NHAP", "GIA NHẬP"])
    offboard_idx = find_col(["NGÀY CHẤM DỨT", "NGAY CHAM DUT", "CHẤM DỨT"])

    if code_idx is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File thiếu cột 'MNV' (Mã cộng tác viên)",
        )

    actor = _actor_name(current_user)
    now_iso = _now()

    # Thu thập mã CTV sẵn có trong bmk_ctv_reconciliations
    existing_reconcil_codes = set()
    async for doc in db[RECONCILIATION_COLLECTION].find({}, {"employeeCode": 1}):
        c = doc.get("employeeCode")
        if c:
            existing_reconcil_codes.add(str(c).strip())

    reconcil_ops = []
    detail_ops = []
    collab_ops = []

    affected_collab_codes = set()
    periods_affected = set()
    new_reconcil_created = 0
    total_valid_rows = 0

    for row in rows[header_row_idx + 1 :]:
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        raw_code = row[code_idx] if code_idx < len(row) else None
        employee_code = parse_employee_code(raw_code)
        if not employee_code:
            continue

        full_name = parse_str_val(row[name_idx]) if (name_idx is not None and name_idx < len(row)) else ""
        tax_code = parse_str_val(row[tax_idx]) if (tax_idx is not None and tax_idx < len(row)) else ""
        acc_name = parse_str_val(row[acc_name_idx]) if (acc_name_idx is not None and acc_name_idx < len(row)) else ""
        acc_num = parse_str_val(row[acc_num_idx]) if (acc_num_idx is not None and acc_num_idx < len(row)) else ""
        bank_name = parse_str_val(row[bank_idx]) if (bank_idx is not None and bank_idx < len(row)) else ""
        taxable_income = parse_float_val(row[taxable_idx]) if (taxable_idx is not None and taxable_idx < len(row)) else 0.0
        tax_tncn = parse_float_val(row[tax_tncn_idx]) if (tax_tncn_idx is not None and tax_tncn_idx < len(row)) else 0.0
        net_income = parse_float_val(row[net_idx]) if (net_idx is not None and net_idx < len(row)) else 0.0
        work_location = parse_str_val(row[work_loc_idx]) if (work_loc_idx is not None and work_loc_idx < len(row)) else ""
        payment_date = parse_date_str(row[pay_date_idx]) if (pay_date_idx is not None and pay_date_idx < len(row)) else None
        payment_content = parse_str_val(row[pay_content_idx]) if (pay_content_idx is not None and pay_content_idx < len(row)) else ""
        onboard_date = parse_date_str(row[onboard_idx]) if (onboard_idx is not None and onboard_idx < len(row)) else None
        offboard_date = parse_date_str(row[offboard_idx]) if (offboard_idx is not None and offboard_idx < len(row)) else None

        if not payment_date:
            # Nếu dòng không có ngày thanh toán hợp lệ thì bỏ qua dòng đó
            continue

        period = calculate_period(payment_date)
        periods_affected.add(period)
        affected_collab_codes.add(employee_code)
        total_valid_rows += 1

        # 1. Kiểm tra tồn tại trong bmk_ctv_reconciliations
        if employee_code not in existing_reconcil_codes:
            contract_expiry_date = add_six_months(onboard_date)
            new_reconcil_doc = {
                "_id": employee_code,
                "employeeCode": employee_code,
                "fullName": full_name or "",
                "idNumber": tax_code or None,
                "contractOnboardDate": onboard_date,
                "contractExpiryDate": contract_expiry_date,
                "isCalculateExpired": True if contract_expiry_date else None,
                "createdSource": "HTR BMK",
                "isBmkSystemExist": False,
                "isSynced": False,
                "departmentLevel1": None,
                "position": None,
                "employmentStatus": "Hiệu lực",
                "onboardDate": onboard_date,
                "offboardDate": offboard_date,
                "tpbankInfo": {"contracts": []},
                "bmkHrInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "bmkSystemInfo": {
                    "contractCount": 0,
                    "idCardCount": 0,
                    "liquidationCount": 0,
                    "taxCommitmentCount": 0,
                },
                "result": None,
                "reconciliationStatus": "pending",
                "createdAt": now_iso,
                "updatedAt": now_iso,
            }
            reconcil_ops.append(
                UpdateOne(
                    {"employeeCode": employee_code},
                    {"$setOnInsert": new_reconcil_doc},
                    upsert=True,
                )
            )
            existing_reconcil_codes.add(employee_code)
            new_reconcil_created += 1

        # 2. Chuẩn bị upsert vào bmk_ctv_expense_details (Unique: employeeCode + paymentDate)
        detail_doc = {
            "employeeCode": employee_code,
            "fullName": full_name,
            "taxCode": tax_code,
            "bankAccount": {
                "accountName": acc_name,
                "accountNumber": acc_num,
                "bankName": bank_name,
            },
            "taxableIncome": taxable_income,
            "personalIncomeTax": tax_tncn,
            "netIncome": net_income,
            "workLocation": work_location,
            "paymentDate": payment_date,
            "paymentContent": payment_content,
            "onboardDate": onboard_date,
            "offboardDate": offboard_date,
            "period": period,
            "updatedAt": now_iso,
            "updatedBy": actor,
        }
        detail_ops.append(
            UpdateOne(
                {"employeeCode": employee_code, "paymentDate": payment_date},
                {
                    "$set": detail_doc,
                    "$setOnInsert": {
                        "createdAt": now_iso,
                        "createdBy": actor,
                    },
                },
                upsert=True,
            )
        )

        # 3. Chuẩn bị upsert thông tin CTV vào bmk_ctv_collaborator_expenses
        collab_doc = {
            "employeeCode": employee_code,
            "fullName": full_name,
            "taxCode": tax_code,
            "latestPaymentDate": payment_date,
            "latestPeriod": period,
            "updatedAt": now_iso,
            "updatedBy": actor,
        }
        collab_ops.append(
            UpdateOne(
                {"employeeCode": employee_code},
                {
                    "$set": collab_doc,
                    "$setOnInsert": {
                        "_id": employee_code,
                        "totalPaymentCount": 0,
                        "createdAt": now_iso,
                        "createdBy": actor,
                    },
                },
                upsert=True,
            )
        )

    # Thực thi bulk operations vào MongoDB
    if reconcil_ops:
        await db[RECONCILIATION_COLLECTION].bulk_write(reconcil_ops, ordered=False)

    imported_details = 0
    updated_details = 0
    if detail_ops:
        res = await db[EXPENSE_DETAILS_COLLECTION].bulk_write(detail_ops, ordered=False)
        imported_details = res.upserted_count
        updated_details = res.modified_count

    if collab_ops:
        await db[COLLABORATOR_EXPENSES_COLLECTION].bulk_write(collab_ops, ordered=False)

    # 4. Tính toán và cập nhật lại chính xác totalPaymentCount cho tất cả CTV có liên quan
    if affected_collab_codes:
        pipeline_counts = [
            {"$match": {"employeeCode": {"$in": list(affected_collab_codes)}}},
            {"$group": {"_id": "$employeeCode", "count": {"$sum": 1}}},
        ]
        counts_res = await db[EXPENSE_DETAILS_COLLECTION].aggregate(pipeline_counts).to_list(len(affected_collab_codes) + 10)
        count_ops = [
            UpdateOne({"employeeCode": item["_id"]}, {"$set": {"totalPaymentCount": item["count"]}})
            for item in counts_res
        ]
        if count_ops:
            await db[COLLABORATOR_EXPENSES_COLLECTION].bulk_write(count_ops, ordered=False)

    logger.info(
        f"Import chi phí CTV hoàn tất bởi {actor}: {total_valid_rows} dòng xử lý, "
        f"{imported_details} bản ghi mới, {updated_details} cập nhật, "
        f"{new_reconcil_created} CTV mới tạo vào bmk_ctv_reconciliations."
    )

    return ExpenseImportResponse(
        success=True,
        totalRows=total_valid_rows,
        importedDetails=imported_details,
        updatedDetails=updated_details,
        newReconciliationCreated=new_reconcil_created,
        collaboratorsCount=len(affected_collab_codes),
        periodsAffected=sorted(list(periods_affected), reverse=True),
        message=f"Import thành công: {total_valid_rows} dòng xử lý ({imported_details} mới, {updated_details} cập nhật), thêm {new_reconcil_created} CTV vào đối soát.",
    )


@router.get("/export")
async def export_expenses(
    period: Optional[str] = Query(None),
    work_location: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    mode: str = Query("details", enum=["details", "collaborators"]),
    db=Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """Xuất file Excel dữ liệu chi phí thanh toán CTV.
    - mode == 'details': Xuất danh sách chi tiết các đợt thanh toán.
    - mode == 'collaborators': Xuất tổng hợp CTV trên từng row bằng template export_thanh_toan_phi_CTV.xlsx,
      các kỳ thanh toán xếp giảm dần theo ngày thanh toán (#1 là lần gần nhất).
    """
    if not isinstance(search, str):
        search = None
    if not isinstance(period, str):
        period = None
    if not isinstance(work_location, str):
        work_location = None

    if mode == "details":
        wb = Workbook()
        ws = wb.active
        ws.title = "Chi tiết thanh toán"
        headers = [
            "STT", "Mã CTV", "Họ và tên", "Mã số thuế", "Tên tài khoản", "Số tài khoản",
            "Ngân hàng", "Tổng TN chịu thuế", "Thuế TNCN", "Thực lĩnh", "Nơi làm việc",
            "Ngày BMK TT", "Nội dung TT", "Ngày gia nhập", "Ngày chấm dứt", "Kỳ thanh toán"
        ]
        ws.append(headers)

        header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        data_font = Font(name="Tahoma", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        query = {}
        if period and period.strip():
            query["period"] = period.strip()
        if work_location and work_location.strip():
            query["workLocation"] = work_location.strip()
        if search and search.strip():
            s = search.strip()
            reg = {"$regex": re.escape(s), "$options": "i"}
            query["$or"] = [
                {"employeeCode": reg},
                {"fullName": reg},
                {"taxCode": reg},
            ]

        cursor = db[EXPENSE_DETAILS_COLLECTION].find(query).sort([("paymentDate", -1), ("employeeCode", 1)])
        row_idx = 2
        async for doc in cursor:
            bank = doc.get("bankAccount") or {}
            row_data = [
                row_idx - 1,
                doc.get("employeeCode", ""),
                doc.get("fullName", ""),
                doc.get("taxCode", ""),
                bank.get("accountName", ""),
                bank.get("accountNumber", ""),
                bank.get("bankName", ""),
                doc.get("taxableIncome", 0.0),
                doc.get("personalIncomeTax", 0.0),
                doc.get("netIncome", 0.0),
                doc.get("workLocation", ""),
                format_display_date(doc.get("paymentDate")),
                doc.get("paymentContent", ""),
                format_display_date(doc.get("onboardDate")),
                format_display_date(doc.get("offboardDate")),
                doc.get("period", ""),
            ]
            ws.append(row_data)

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if col_idx in (1, 2, 4, 6, 12, 14, 15, 16):
                    cell.alignment = align_center
                elif col_idx in (8, 9, 10):
                    cell.alignment = align_right
                    cell.number_format = "#,##0"
                else:
                    cell.alignment = align_left

            row_idx += 1

        filename = f"chi_tiet_chi_phi_ctv_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    else:
        # Sử dụng template export_thanh_toan_phi_CTV.xlsx theo yêu cầu
        if not EXPORT_TEMPLATE_PATH.exists():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="File mẫu export_thanh_toan_phi_CTV.xlsx không tồn tại trên hệ thống",
            )

        from openpyxl.utils import get_column_letter

        wb = load_workbook(str(EXPORT_TEMPLATE_PATH))
        sheet_name = "Gui BMK" if "Gui BMK" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # Xóa các dòng placeholder cũ (từ dòng 10 trở đi)
        if ws.max_row >= 10:
            ws.delete_rows(10, ws.max_row - 9)

        query = {}
        if search and search.strip():
            s = search.strip()
            reg = {"$regex": re.escape(s), "$options": "i"}
            query["$or"] = [
                {"employeeCode": reg},
                {"fullName": reg},
                {"taxCode": reg},
            ]

        collabs = (
            await db[COLLABORATOR_EXPENSES_COLLECTION]
            .find(query)
            .sort([("totalPaymentCount", -1), ("employeeCode", 1)])
            .to_list(10000)
        )
        collab_codes = [c.get("employeeCode") for c in collabs if c.get("employeeCode")]

        # Lấy tất cả chi tiết thanh toán của các CTV này, sắp xếp giảm dần theo ngày thanh toán
        details_cursor = db[EXPENSE_DETAILS_COLLECTION].find(
            {"employeeCode": {"$in": collab_codes}}
        ).sort([("paymentDate", -1), ("_id", -1)])
        all_details = await details_cursor.to_list(100000)

        payments_by_code = defaultdict(list)
        for d in all_details:
            c_code = d.get("employeeCode")
            if c_code:
                payments_by_code[c_code].append(d)

        data_font = Font(name="Times New Roman", size=11)
        thin_border = Border(
            left=Side(style="thin", color="000000"),
            right=Side(style="thin", color="000000"),
            top=Side(style="thin", color="000000"),
            bottom=Side(style="thin", color="000000"),
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        num_format_currency = '#,##0'

        # Xác định số kỳ thanh toán tối đa để mở rộng cột tiêu đề nếu > 12 kỳ
        max_payments = max([len(payments_by_code.get(c, [])) for c in collab_codes] or [1])
        if max_payments > 12:
            h_font = Font(name="Times New Roman", size=11, bold=True)
            h_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
            for p_i in range(12, max_payments):
                b_col = 7 + p_i * 4
                ws.merge_cells(start_row=8, start_column=b_col, end_row=8, end_column=b_col + 3)
                c8 = ws.cell(row=8, column=b_col, value=f"#{p_i + 1}")
                c8.font = h_font
                c8.alignment = align_center
                c8.fill = h_fill
                for offset in range(4):
                    ws.cell(row=8, column=b_col + offset).border = thin_border

                sub_titles = ["Tổng TN \nchịu thuế", "THUẾ TNCN", "THỰC LĨNH", "NGÀY BMK TT"]
                sub_widths = [18.9, 17.3, 17.0, 15.5]
                for offset, (title, width) in enumerate(zip(sub_titles, sub_widths)):
                    sub_col = b_col + offset
                    c9 = ws.cell(row=9, column=sub_col, value=title)
                    c9.font = h_font
                    c9.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c9.fill = h_fill
                    c9.border = thin_border
                    col_letter = get_column_letter(sub_col)
                    ws.column_dimensions[col_letter].width = width

        total_cols = max(54, 6 + max_payments * 4)

        for idx, collab in enumerate(collabs, start=1):
            r_num = 9 + idx
            code = collab.get("employeeCode", "")
            full_name = collab.get("fullName", "")
            tax_code = collab.get("taxCode", "")
            ctv_payments = payments_by_code.get(code, [])
            total_payments_count = len(ctv_payments) or collab.get("totalPaymentCount", 0)
            total_taxable_income = sum(p.get("taxableIncome", 0.0) for p in ctv_payments)

            # Cột thông tin cơ bản
            c_stt = ws.cell(row=r_num, column=1, value=idx)
            c_code = ws.cell(row=r_num, column=2, value=code)
            c_name = ws.cell(row=r_num, column=3, value=full_name)
            c_tax = ws.cell(row=r_num, column=4, value=tax_code)
            c_count = ws.cell(row=r_num, column=5, value=total_payments_count)
            c_tot_inc = ws.cell(row=r_num, column=6, value=total_taxable_income)

            c_stt.alignment = align_center
            c_code.alignment = align_center
            c_name.alignment = align_left
            c_tax.alignment = align_center
            c_count.alignment = align_center
            c_tot_inc.alignment = align_right
            c_tot_inc.number_format = num_format_currency

            for cell_obj in (c_stt, c_code, c_name, c_tax, c_count, c_tot_inc):
                cell_obj.font = data_font
                cell_obj.border = thin_border

            # Điền các kỳ thanh toán giảm dần theo ngày thanh toán
            for p_idx, p in enumerate(ctv_payments):
                b_col = 7 + p_idx * 4
                c_taxable = ws.cell(row=r_num, column=b_col, value=p.get("taxableIncome", 0.0))
                c_tncn = ws.cell(row=r_num, column=b_col + 1, value=p.get("personalIncomeTax", 0.0))
                c_net = ws.cell(row=r_num, column=b_col + 2, value=p.get("netIncome", 0.0))
                c_date = ws.cell(row=r_num, column=b_col + 3, value=format_display_date(p.get("paymentDate")))

                for num_cell in (c_taxable, c_tncn, c_net):
                    num_cell.alignment = align_right
                    num_cell.number_format = num_format_currency
                    num_cell.font = data_font
                    num_cell.border = thin_border

                c_date.alignment = align_center
                c_date.font = data_font
                c_date.border = thin_border

            # Viền border cho các cột còn lại nếu CTV có ít hơn số kỳ tối đa
            filled_cols = 6 + len(ctv_payments) * 4
            for empty_col in range(filled_cols + 1, total_cols + 1):
                empty_cell = ws.cell(row=r_num, column=empty_col)
                empty_cell.border = thin_border

        filename = f"tong_hop_chi_phi_ctv_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    output = BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
